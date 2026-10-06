"""The bundled ``xlsx`` skill: its scripts run through ``valuz-python`` and ``dsoffice``.

Every script is invoked exactly as a session would (``valuz-python <skill>/scripts/x.py``)
on workbooks created in ``tmp_path``. Assertions check real outcomes: error locations
after a LibreOffice recalculation, recalculated values, audit findings, rendered PNGs,
the structural checker and a PDF conversion of the files the scripts produced.
"""

from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.workbook.defined_name import DefinedName
from PIL import Image

from tests.resources.office_skills.conftest import SKILLS_DIR, run

SKILL = SKILLS_DIR / "xlsx"
SCRIPTS = SKILL / "scripts"
DEEPSEEK_CHECKER = (
    Path(__file__).resolve().parents[3]
    / "vendor/dsh-runtime/node_modules/@deepseek-ai/dsh-skill-office/assets/scripts/check_office.py"
)


def script(name: str, *args: object, env: dict[str, str], cwd: Path) -> subprocess.CompletedProcess:
    """Run a skill script with valuz-python without raising on a non-zero exit."""
    argv = ["valuz-python", str(SCRIPTS / name), *map(str, args)]
    return subprocess.run(argv, env=env, cwd=cwd, capture_output=True, text=True)


def report_of(completed: subprocess.CompletedProcess) -> dict:
    assert completed.stdout, completed.stderr
    return json.loads(completed.stdout)


def locations(report: dict) -> dict[str, list[str]]:
    return {name: group["locations"] for name, group in report["errors_by_type"].items()}


def assert_checker_and_pdf(path: Path, env: dict[str, str], cwd: Path) -> None:
    checked = report_of(
        run(["valuz-python", str(SCRIPTS / "check_office.py"), str(path)], env, cwd)
    )
    assert checked["verdict"] == "pass"
    pdf = cwd / f"{path.stem}-check.pdf"
    run(["dsoffice", "convert", "--input", str(path), "--output", str(pdf)], env, cwd)
    assert pdf.read_bytes().startswith(b"%PDF")


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


# ---------------------------------------------------------------------------
# Skill layout


def test_skill_files_and_frontmatter() -> None:
    text = (SKILL / "SKILL.md").read_text(encoding="utf-8")
    frontmatter = text.split("---", 2)[1]
    assert re.search(r"^name: xlsx$", frontmatter, re.M)
    assert re.search(r"^description: \S.{100,}$", frontmatter, re.M)
    for name in ("check_office.py", "recalc_check.py", "model_audit.py"):
        assert (SCRIPTS / name).is_file()
    for link in re.findall(r"`(reference/[\w.-]+\.md)`", text):
        assert (SKILL / link).is_file(), link
    every_doc = [text, *(path.read_text(encoding="utf-8") for path in SKILL.glob("reference/*.md"))]
    for document in every_doc:
        assert "../" not in document
        for stale in ("load_workspace_dependencies", "read_image", "render_document", "<node>"):
            assert stale not in document
        assert "present(" not in document
    if DEEPSEEK_CHECKER.is_file():
        assert (SCRIPTS / "check_office.py").read_bytes() == DEEPSEEK_CHECKER.read_bytes()


# ---------------------------------------------------------------------------
# recalc_check.py


def _error_workbook(path: Path) -> None:
    wb = Workbook()
    calc = wb.active
    calc.title = "Calc"
    calc["A1"] = 10
    calc["A2"] = 0
    calc["B1"] = "=A1/A2"  # #DIV/0!
    calc["B2"] = "=INDEX(A1:A2,5)"  # #REF!
    calc["B3"] = "=NOSUCHFUNC(A1)"  # #NAME?
    calc["B4"] = '=A1+"abc"'  # #VALUE!
    calc["B5"] = "=VLOOKUP(99,A1:A2,1,FALSE)"  # #N/A
    calc["B6"] = "=10^400"  # #NUM!
    calc["B7"] = "=A1 A2"  # #NULL!
    calc["B8"] = "=SUM(A1:A2)"  # fine: 10
    calc["B9"] = "=XLOOKUP(10,A1:A2,A1:A2)"  # #NAME?: Excel 2010+ function without _xlfn.
    calc["B10"] = "=_xlfn.XLOOKUP(10,A1:A2,A1:A2)"  # fine: 10
    calc["C1"] = "=B1*2"  # #DIV/0! inherited from B1
    other = wb.create_sheet("Other Sheet")
    other["A1"] = "=Calc!B2+1"  # #REF! inherited across sheets
    other["A2"] = "=Calc!B8*3"  # fine: 30
    wb.save(path)


def test_recalc_check_reports_every_error_type(
    office_env: dict[str, str], needs_dsoffice: None, tmp_path: Path
) -> None:
    source = tmp_path / "errors.xlsx"
    _error_workbook(source)
    before = digest(source)

    completed = script("recalc_check.py", source, env=office_env, cwd=tmp_path)
    assert completed.returncode == 1, completed.stderr
    report = report_of(completed)
    assert report["status"] == "errors"
    assert report["total_formulas"] == 13
    assert report["total_errors"] == 10
    assert locations(report) == {
        "#DIV/0!": ["Calc!B1", "Calc!C1"],
        "#REF!": ["Calc!B2", "'Other Sheet'!A1"],
        "#NAME?": ["Calc!B3", "Calc!B9"],
        "#VALUE!": ["Calc!B4"],
        "#N/A": ["Calc!B5"],
        "#NUM!": ["Calc!B6"],
        "#NULL!": ["Calc!B7"],
    }
    assert all(group["hint"] for group in report["errors_by_type"].values())
    roots = {cause["cell"]: cause for cause in report["root_causes"]}
    assert set(roots) == {f"Calc!B{row}" for row in (1, 2, 3, 4, 5, 6, 7, 9)}
    assert roots["Calc!B1"]["formula"] == "=A1/A2"
    assert {cause["reason"] for cause in roots.values()} == {"origin"}
    assert report["sheets"] == [
        {"name": "Calc", "formulas": 11, "errors": 9},
        {"name": "Other Sheet", "formulas": 2, "errors": 1},
    ]

    # The input is untouched; the recalculated copy is a new file with real values.
    assert digest(source) == before
    output = Path(report["output"])
    assert output == tmp_path / "errors.recalc.xlsx"
    values = load_workbook(output, data_only=True)
    assert values["Calc"]["B8"].value == 10
    assert values["Calc"]["B10"].value == 10
    assert values["Other Sheet"]["A2"].value == 30
    assert load_workbook(output)["Calc"]["B8"].value == "=SUM(A1:A2)"
    assert_checker_and_pdf(output, office_env, tmp_path)

    # --allow-errors keeps the report but exits 0; the next default name is used.
    tolerated = script("recalc_check.py", source, "--allow-errors", env=office_env, cwd=tmp_path)
    assert tolerated.returncode == 0
    again = report_of(tolerated)
    assert again["total_errors"] == 10
    assert again["output"] == str(tmp_path / "errors.recalc-2.xlsx")


def test_recalc_check_clean_workbook_values(
    office_env: dict[str, str], needs_dsoffice: None, tmp_path: Path
) -> None:
    wb = Workbook()
    data = wb.active
    data.title = "Data"
    data.append(["Region", "Product", "Sales"])
    for row in [("East", "A", 120), ("East", "B", 80), ("West", "A", 200), ("West", "B", 0)]:
        data.append(row)
    data["E1"] = "Growth"
    data["F1"] = 0.1
    summary = wb.create_sheet("汇总 Summary")
    summary["A1"] = "East total"
    summary["B1"] = '=SUMIFS(Data!C2:C5,Data!A2:A5,"East")'
    summary["A2"] = "West next year"
    summary["B2"] = '=SUMIFS(Data!C2:C5,Data!A2:A5,"West")*(1+Data!F1)'
    summary["A3"] = "Share of B in West"
    summary["B3"] = "=IF(B2=0,0,Data!C5/B2)"
    summary["A4"] = "Lookup West A"
    summary["B4"] = '=_xlfn.XLOOKUP("West",Data!A2:A5,Data!C2:C5)'
    summary["A5"] = "Average"
    summary["B5"] = "=AVERAGE(Data!C2:C5)"
    draft = tmp_path / "summary-draft.xlsx"
    wb.save(draft)

    final = tmp_path / "summary.xlsx"
    completed = script("recalc_check.py", draft, "--output", final, env=office_env, cwd=tmp_path)
    assert completed.returncode == 0, completed.stdout + completed.stderr
    report = report_of(completed)
    assert report["status"] == "clean"
    assert report["total_formulas"] == 5
    assert report["total_errors"] == 0
    assert report["root_causes"] == []
    assert report["fidelity"]["dropped"] == []
    assert any("file to deliver" in note for note in report["notes"])

    values = load_workbook(final, data_only=True)["汇总 Summary"]
    assert values["B1"].value == 200
    assert values["B2"].value == pytest.approx(220)
    assert values["B3"].value == 0
    assert values["B4"].value == 200
    assert values["B5"].value == 100
    formulas = load_workbook(final)["汇总 Summary"]
    assert formulas["B1"].value == '=SUMIFS(Data!C2:C5,Data!A2:A5,"East")'
    assert_checker_and_pdf(final, office_env, tmp_path)

    # After a fix, --replace swaps a fresh copy in place of the previous output.
    summary["B5"] = "=MAX(Data!C2:C5)"
    wb.save(draft)
    again = ["recalc_check.py", draft, "--output", final]
    assert script(*again, env=office_env, cwd=tmp_path).returncode == 2
    replaced = script(*again, "--replace", env=office_env, cwd=tmp_path)
    assert replaced.returncode == 0, replaced.stdout + replaced.stderr
    assert load_workbook(final, data_only=True)["汇总 Summary"]["B5"].value == 200
    assert sorted(path.name for path in tmp_path.glob("*.xlsx")) == [
        "summary-draft.xlsx",
        "summary.xlsx",
    ]
    same = script(
        "recalc_check.py", draft, "--output", draft, "--replace", env=office_env, cwd=tmp_path
    )
    assert same.returncode == 2  # never the input


def test_recalc_check_circular_external_and_names(
    office_env: dict[str, str], needs_dsoffice: None, tmp_path: Path
) -> None:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "S"
    sheet["A1"] = 100
    sheet["A2"] = "=A1+A3*0.1"  # A2 <-> A3: circular
    sheet["A3"] = "=A2*0.5"
    sheet["A4"] = "=A2*2"  # downstream of the loop
    sheet["B1"] = 0
    sheet["B2"] = "=A1/B1"
    sheet["B3"] = "=rate*2"  # reads B2 through a defined name
    sheet["C1"] = "#N/A"
    sheet["C1"].data_type = "s"  # text that only looks like an error
    sheet["D1"] = "='[Budget.xlsx]Sheet1'!A1"
    wb.defined_names["rate"] = DefinedName("rate", attr_text="S!$B$2")
    wb.defined_names["gone"] = DefinedName("gone", attr_text="S!#REF!")
    source = tmp_path / "loops.xlsx"
    wb.save(source)

    completed = script("recalc_check.py", source, env=office_env, cwd=tmp_path)
    assert completed.returncode == 1
    report = report_of(completed)
    assert {item["cell"] for item in report["circular_references"]} == {"S!A2", "S!A3"}
    reasons = {cause["cell"]: cause["reason"] for cause in report["root_causes"]}
    assert reasons == {
        "S!A2": "circular",
        "S!A3": "circular",
        "S!B2": "origin",
        "S!D1": "external-link",
    }
    assert "S!A4" in locations(report)["#VALUE!"]
    assert "S!B3" in locations(report)["#DIV/0!"]
    assert report["external_reference_cells"]["locations"] == ["S!D1"]
    assert report["error_like_text"]["locations"] == ["S!C1"]
    assert report["broken_defined_names"] == ["gone"]
    assert any("other workbooks" in note for note in report["notes"])


def _handmade_workbook(path: Path) -> None:
    """A package written by hand: shared formulas, cells without r, inline strings."""
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    parts = {
        "[Content_Types].xml": (
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.'
            'relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-'
            'officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/'
            'sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.'
            'worksheet+xml"/></Types>'
        ),
        "_rels/.rels": (
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="rId1" Type="{rel}/officeDocument" Target="xl/workbook.xml"/>'
            "</Relationships>"
        ),
        "xl/workbook.xml": (
            f'<workbook xmlns="{main}" xmlns:r="{rel}"><sheets>'
            '<sheet name="Data" sheetId="1" r:id="rId1"/></sheets></workbook>'
        ),
        "xl/_rels/workbook.xml.rels": (
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="rId1" Type="{rel}/worksheet" Target="worksheets/sheet1.xml"/>'
            "</Relationships>"
        ),
        "xl/worksheets/sheet1.xml": (
            f'<worksheet xmlns="{main}"><sheetData>'
            '<row r="1"><c r="A1"><v>0</v></c><c r="B1"><v>4</v></c>'
            '<c r="C1" t="e"><f t="shared" ref="C1:C3" si="0">B1/A1</f><v>#DIV/0!</v></c></row>'
            '<row r="2"><c r="A2"><v>2</v></c><c r="B2"><v>4</v></c>'
            '<c r="C2"><f t="shared" si="0"/><v>2</v></c></row>'
            '<row r="3"><c r="A3"><v>0</v></c><c r="B3"><v>4</v></c>'
            '<c r="C3" t="e"><f t="shared" si="0"/><v>#DIV/0!</v></c></row>'
            '<row r="4"><c><v>1</v></c><c t="inlineStr"><is><t>#N/A</t></is></c>'
            '<c t="str"><f>IF(A4&gt;0,"","x")</f><v></v></c></row>'
            "</sheetData></worksheet>"
        ),
    }
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, xml in parts.items():
            archive.writestr(name, xml)


def test_recalc_check_scans_saved_values_without_recalculating(
    office_env: dict[str, str], needs_dsoffice: None, tmp_path: Path
) -> None:
    handmade = tmp_path / "handmade.xlsx"
    _handmade_workbook(handmade)
    scanned = script("recalc_check.py", handmade, "--no-recalc", env=office_env, cwd=tmp_path)
    assert scanned.returncode == 1
    report = report_of(scanned)
    assert report["recalculated"] is False
    assert report["total_formulas"] == 4
    assert locations(report) == {"#DIV/0!": ["Data!C1", "Data!C3"]}
    formulas = {cause["cell"]: cause["formula"] for cause in report["root_causes"]}
    assert formulas == {"Data!C1": "=B1/A1", "Data!C3": "=B3/A3"}  # shared formula translated
    assert report["error_like_text"]["locations"] == ["Data!B4"]
    assert report["cached_values_missing"]["count"] == 0  # an empty cached string is a result

    # A file written by openpyxl stores no results: the saved-value scan cannot vouch for it.
    wb = Workbook()
    wb.active["A1"] = 2
    wb.active["A2"] = "=A1*3"
    fresh = tmp_path / "fresh.xlsx"
    wb.save(fresh)
    unverified = script("recalc_check.py", fresh, "--no-recalc", env=office_env, cwd=tmp_path)
    assert unverified.returncode == 1
    assert report_of(unverified)["status"] == "unverified"
    recalculated = script("recalc_check.py", fresh, env=office_env, cwd=tmp_path)
    assert recalculated.returncode == 0
    output = report_of(recalculated)["output"]
    rescanned = script("recalc_check.py", output, "--no-recalc", env=office_env, cwd=tmp_path)
    assert rescanned.returncode == 0
    assert report_of(rescanned)["status"] == "clean"


def test_recalc_check_reports_features_lost_in_the_round_trip(
    office_env: dict[str, str], needs_dsoffice: None, tmp_path: Path
) -> None:
    wb = Workbook()
    wb.active.title = "S"
    wb.active["A1"] = 2
    wb.active["B2"] = "=A1*3"
    plain = tmp_path / "plain.xlsx"
    wb.save(plain)
    threaded = tmp_path / "threaded.xlsx"
    with zipfile.ZipFile(plain) as source, zipfile.ZipFile(threaded, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"</Types>",
                    b'<Override PartName="/xl/threadedComments/threadedComment1.xml" '
                    b'ContentType="application/vnd.ms-excel.threadedcomments+xml"/></Types>',
                )
            target.writestr(item, data)
        target.writestr(
            "xl/worksheets/_rels/sheet1.xml.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.microsoft.com/office/2017/10/'
            'relationships/threadedComment" Target="../threadedComments/threadedComment1.xml"/>'
            "</Relationships>",
        )
        target.writestr(
            "xl/threadedComments/threadedComment1.xml",
            '<ThreadedComments xmlns="http://schemas.microsoft.com/office/spreadsheetml/2018/'
            'threadedcomments"><threadedComment ref="B2" dT="2025-01-01T00:00:00.00" '
            'personId="{11111111-1111-1111-1111-111111111111}" '
            'id="{22222222-2222-2222-2222-222222222222}"><text>Check</text>'
            "</threadedComment></ThreadedComments>",
        )
    completed = script("recalc_check.py", threaded, env=office_env, cwd=tmp_path)
    assert completed.returncode == 0
    report = report_of(completed)
    assert report["fidelity"]["dropped"] == [
        {"feature": "threaded_comments", "before": 1, "after": 0}
    ]
    assert any("Deliver the source workbook" in note for note in report["notes"])


@pytest.mark.skipif(os.name == "nt", reason="builds a POSIX PATH with a symlinked wrapper")
def test_recalc_check_refusals(office_env: dict[str, str], tmp_path: Path) -> None:
    wb = Workbook()
    wb.active["A1"] = "=1+1"
    source = tmp_path / "book.xlsx"
    wb.save(source)
    taken = tmp_path / "taken.xlsx"
    taken.write_bytes(b"keep me")
    refused = script("recalc_check.py", source, "--output", taken, env=office_env, cwd=tmp_path)
    assert refused.returncode == 2
    assert "already exists" in refused.stderr
    assert taken.read_bytes() == b"keep me"

    # A session without dsoffice gets a JSON failure report and exit 3.
    lonely_bin = tmp_path / "bin"
    lonely_bin.mkdir()
    python = shutil.which("valuz-python", path=office_env["PATH"])
    assert python is not None
    (lonely_bin / "valuz-python").symlink_to(python)
    env = {**office_env, "PATH": f"{lonely_bin}:/usr/bin:/bin"}
    missing = script("recalc_check.py", source, env=env, cwd=tmp_path)
    assert missing.returncode == 3
    failure = report_of(missing)
    assert failure["status"] == "failed"
    assert "dsoffice" in failure["detail"]


# ---------------------------------------------------------------------------
# Financial-model conventions: the worked example and model_audit.py


def _worked_example_code() -> str:
    text = (SKILL / "reference" / "financial-models.md").read_text(encoding="utf-8")
    section = text.split("## 6. Worked example", 1)[1]
    match = re.search(r"```python\n(.*?)```", section, re.S)
    assert match is not None
    return match.group(1)


def test_worked_example_applies_the_conventions(
    office_env: dict[str, str], needs_dsoffice: None, tmp_path: Path
) -> None:
    (tmp_path / "example.py").write_text(_worked_example_code(), encoding="utf-8")
    run(["valuz-python", "example.py"], office_env, tmp_path)

    checked = script("recalc_check.py", "projection.xlsx", env=office_env, cwd=tmp_path)
    assert checked.returncode == 0, checked.stdout
    report = report_of(checked)
    assert report["status"] == "clean"
    assert report["total_formulas"] == 27
    output = Path(report["output"])

    values = load_workbook(output, data_only=True)["Model"]
    assert values["C4"].value == pytest.approx(1350)
    assert values["G4"].value == pytest.approx(1250 * 1.08**5)
    assert values["C6"].value == pytest.approx(297)
    assert values["C10"].value == pytest.approx(2524.5)
    assert values["C12"].value == pytest.approx(2204.5)

    # Styling as the example wrote it (LibreOffice re-escapes number formats on save).
    source = load_workbook(tmp_path / "projection.xlsx")
    inputs, model = source["Inputs"], source["Model"]
    assert inputs["B2"].font.color.rgb.endswith("0000FF")  # typed input: blue
    assert inputs["B3"].fill.fgColor.rgb.endswith("FFF2CC")  # key assumption
    assert model["B4"].font.color.rgb.endswith("008000")  # link to another sheet: green
    assert model["C4"].value == "=B4*(1+C3)"
    assert model["C4"].font.color.rgb.endswith("000000")  # same-sheet formula: black
    assert all(isinstance(model[f"{col}2"].value, str) for col in "BCDEFG")  # years as text
    assert model["C4"].number_format == '#,##0.0_);(#,##0.0);"-"_)'
    assert model["C3"].number_format == '0.0%_);(0.0%);"-"_)'
    assert model["C9"].number_format == '0.0"x"_)'
    # The delivered (recalculated) copy keeps one font and the colours.
    delivered = load_workbook(output)
    cells = [cell for ws in delivered for row in ws.iter_rows() for cell in row if cell.value]
    assert {cell.font.name for cell in cells} == {"Arial"}
    assert delivered["Inputs"]["B2"].font.color.rgb.endswith("0000FF")
    assert delivered["Model"]["B4"].font.color.rgb.endswith("008000")

    audited = script("model_audit.py", output, "--strict", env=office_env, cwd=tmp_path)
    assert audited.returncode == 0, audited.stdout
    audit = report_of(audited)
    assert audit["finding_count"] == 0
    assert audit["fonts"]["consistent"] is True

    preview = tmp_path / "preview-model"
    run(
        ["dsoffice", "render", "--input", str(output), "--output-dir", str(preview)]
        + ["--sheet", "Model", "--range", "A1:G12", "--dpi", "96"],
        office_env,
        tmp_path,
    )
    manifest = json.loads((preview / "manifest.json").read_text(encoding="utf-8"))
    image = Path(manifest["images"][0]["path"])
    assert image.is_file()
    with Image.open(image) as png:
        rgba = png.convert("RGBA")
        flat = Image.new("RGB", rgba.size, "white")
        flat.paste(rgba, mask=rgba.split()[3])
        darkest, lightest = flat.convert("L").getextrema()
        assert darkest < 100 < lightest  # text on a light background, not a blank tile
    assert_checker_and_pdf(output, office_env, tmp_path)


def test_model_audit_flags_convention_breaks(office_env: dict[str, str], tmp_path: Path) -> None:
    blue, green, black = (Font(name="Arial", color=c) for c in ("0000FF", "008000", "000000"))
    wb = Workbook()
    inputs = wb.active
    inputs.title = "Inputs"
    for ref, value, font in (("B1", 0.08, blue), ("B2", 0.3, black)):  # B2: input not blue
        inputs[ref] = value
        inputs[ref].font = font
    model = wb.create_sheet("Model")
    model["B2"] = 100
    model["B2"].font = blue
    for col in "CDF":
        model[f"{col}2"] = f"={chr(ord(col) - 1)}2*(1+Inputs!$B$1)"
        model[f"{col}2"].font = green
    model["E2"] = 140  # a pasted value between formulas
    model["E2"].font = blue
    for col in "BCDEF":
        model[f"{col}3"] = f"={col}2*Inputs!$B$2"
        model[f"{col}3"].font = green
    model["D3"] = "=D2*0.3"  # hard-coded and off-pattern
    model["D3"].font = Font(name="Calibri")  # and a second font family
    model["B4"] = "=1250+380"  # constants only, coloured like an input
    model["B4"].font = blue
    model["C4"] = "='[Budget.xlsx]Sheet1'!A1"  # external link not red
    model["C4"].font = black
    path = tmp_path / "model.xlsx"
    wb.save(path)

    advisory = script("model_audit.py", path, env=office_env, cwd=tmp_path)
    assert advisory.returncode == 0
    report = report_of(advisory)
    mismatches = {item["cell"]: item for item in report["colour_coding"]["mismatches"]["cells"]}
    assert set(mismatches) == {"Inputs!B2", "Model!B4", "Model!C4"}
    assert mismatches["Inputs!B2"]["category"] == "input"
    assert mismatches["Model!B4"]["expected"] == "black"
    assert mismatches["Model!C4"]["category"] == "external"
    hardcoded = {item["cell"]: item["numbers"] for item in report["hardcoded_numbers"]["cells"]}
    assert hardcoded == {"Model!D3": ["0.3"], "Model!B4": ["1250", "380"]}
    inconsistent = report["inconsistent_formulas"]["cells"]
    assert [item["cell"] for item in inconsistent] == ["Model!D3"]
    assert inconsistent[0]["pattern_formula"] == "=B2*Inputs!$B$2"
    assert [item["cell"] for item in report["constants_in_formula_runs"]["cells"]] == ["Model!E2"]
    assert report["fonts"]["families"] == {"Arial": 13, "Calibri": 1}
    assert report["fonts"]["outliers"]["cells"] == ["Model!D3"]
    assert report["finding_count"] == 8

    strict = script("model_audit.py", path, "--strict", env=office_env, cwd=tmp_path)
    assert strict.returncode == 1
    limited = report_of(
        script("model_audit.py", path, "--sheets", "Inputs", env=office_env, cwd=tmp_path)
    )
    assert limited["sheets"] == ["Inputs"]
    assert limited["finding_count"] == 1


def test_model_audit_recognises_an_existing_scheme(
    office_env: dict[str, str], tmp_path: Path
) -> None:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Budget"
    for row in range(1, 7):  # inputs typed in black throughout: the file's own scheme
        sheet.cell(row=row, column=1, value=row * 10).font = Font(name="Arial")
        sheet.cell(row=row, column=2, value=f"=A{row}*2").font = Font(name="Arial")
    path = tmp_path / "budget.xlsx"
    wb.save(path)
    report = report_of(script("model_audit.py", path, env=office_env, cwd=tmp_path))
    assert report["colour_coding"]["observed"]["input"] == {"black": 6}
    assert any("input cells are mostly black" in note for note in report["notes"])
    column = report_of(
        script("model_audit.py", path, "--axis", "column", env=office_env, cwd=tmp_path)
    )
    assert column["inconsistent_formulas"]["count"] == 0
