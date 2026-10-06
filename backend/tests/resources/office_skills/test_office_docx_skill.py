"""The bundled ``docx`` skill: its scripts run through ``valuz-python`` and ``dsoffice``.

Every script is invoked as a session would (``valuz-python <skill>/scripts/x.py``) on
documents built in ``tmp_path`` by build scripts that import ``docx_helpers``. Assertions
read the produced XML (revisions, comments, fields, numbering, sections), the scripts'
JSON, the structural checker, and LibreOffice output (PDF conversion, text export, PNGs).
"""

from __future__ import annotations

import json
import re
import subprocess
import textwrap
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest

from tests.resources.office_skills.conftest import SKILLS_DIR, run

SKILL = SKILLS_DIR / "docx"
SCRIPTS = SKILL / "scripts"
DEEPSEEK_CHECKER = (
    Path(__file__).resolve().parents[3]
    / "vendor/dsh-runtime/node_modules/@deepseek-ai/dsh-skill-office/assets/scripts/check_office.py"
)
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
DATE = "2026-10-04T09:30:00Z"


def script(name: str, *args: object, env: dict[str, str], cwd: Path) -> subprocess.CompletedProcess:
    """Run a skill script with valuz-python without raising on a non-zero exit."""
    argv = ["valuz-python", str(SCRIPTS / name), *map(str, args)]
    return subprocess.run(argv, env=env, cwd=cwd, capture_output=True, text=True)


def build(source: str, env: dict[str, str], cwd: Path) -> None:
    """Run a python-docx build script that imports the skill's helpers."""
    path = cwd / "build.py"
    header = f"import sys\nsys.path.insert(0, {str(SCRIPTS)!r})\n"
    path.write_text(header + textwrap.dedent(source), encoding="utf-8")
    run(["valuz-python", str(path)], env, cwd)


def part(path: Path, name: str) -> ET.Element:
    with zipfile.ZipFile(path) as archive:
        return ET.fromstring(archive.read(name))


def members(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as archive:
        return archive.namelist()


def paragraph_texts(path: Path) -> list[str]:
    body = part(path, "word/document.xml").find(f"{W}body")
    return ["".join(t.text or "" for t in p.iter(f"{W}t")) for p in body.iter(f"{W}p")]


def assert_checker_passes(path: Path, env: dict[str, str], *contains: str) -> dict:
    args = [arg for text in contains for arg in ("--contains", text)]
    result = script("check_office.py", path, *args, env=env, cwd=path.parent)
    report = json.loads(result.stdout)
    assert result.returncode == 0 and report["verdict"] == "pass", result.stdout
    return report


def assert_converts_to_pdf(path: Path, env: dict[str, str]) -> None:
    pdf = path.with_name(path.stem + "-check.pdf")
    run(["dsoffice", "convert", "--input", str(path), "--output", str(pdf)], env, path.parent)
    assert pdf.read_bytes().startswith(b"%PDF")


# ---------------------------------------------------------------- layout


def test_skill_layout_is_self_contained():
    text = (SKILL / "SKILL.md").read_text(encoding="utf-8")
    assert text.startswith("---\nname: docx\ndescription: ")
    assert len(text.splitlines()) <= 300
    for doc in [SKILL / "SKILL.md", *sorted((SKILL / "reference").glob("*.md"))]:
        body = doc.read_text(encoding="utf-8")
        assert "../" not in body, f"{doc.name} reaches outside the skill"
        for removed in (
            "load_workspace_dependencies",
            "read_image",
            "render_document",
            "<node>",
            "present(",
        ):
            assert removed not in body, f"{doc.name} still mentions {removed}"
        for link in re.findall(r"\]\((reference/[^)]+)\)", body):
            assert (SKILL / link).is_file(), link
    for name in (
        "docx_inspect.py",
        "docx_review.py",
        "docx_pack.py",
        "docx_toc.py",
        "docx_helpers.py",
    ):
        assert (SCRIPTS / name).is_file()
    if DEEPSEEK_CHECKER.is_file():
        assert (SCRIPTS / "check_office.py").read_bytes() == DEEPSEEK_CHECKER.read_bytes()


# ---------------------------------------------------------------- creation helpers

RICH_DOC = """
from docx import Document
from docx.shared import Cm, Pt
from PIL import Image
from docx_helpers import *

Image.new("RGB", (2400, 1200), (31, 56, 100)).save("wide.png", dpi=(100, 100))
doc = Document()
section = doc.sections[0]
set_page_layout(section, "A4", margins=Cm(2))
set_document_fonts(doc, latin="Calibri", east_asia="SimSun", size=11, east_asia_lang="zh-CN")
set_style_font(doc, "Heading 1", latin="Arial", east_asia="SimHei", size=16, color="1F3864")
add_page_number_footer(section, "Page {PAGE} of {NUMPAGES}")
doc.add_heading("年度报告 Annual report", 1)
p = doc.add_paragraph("收入增长 12%")
add_footnote(p, "数据来源：公司年报。")
p.add_run(" see ")
add_hyperlink(p, "the filing", "https://example.com/filing")
add_endnote(p, "Method note.")
bullets = create_list(doc, "bullet")
add_list_item(doc, "要点一 First", bullets)
add_list_item(doc, "Second level", bullets, level=1)
steps = create_list(doc, "number")
add_list_item(doc, "Step one", steps)
add_list_item(doc, "Step two", steps)
restart = create_list(doc, "number")
add_list_item(doc, "Restarted", restart)
rows = [["Region", "Revenue"]] + [[f"R{i}", f"{i * 100:,}"] for i in range(1, 6)]
t = add_table(doc, rows, col_widths=[Cm(8), Cm(4)], align=["left", "right"])
t.cell(1, 0).merge(t.cell(1, 1))
shade_cell(t.cell(2, 0), "FFF2CC")
set_cell_borders(t.cell(3, 0), bottom={"sz": 12, "color": "C00000"})
add_picture_fit(doc, "wide.png")
p = doc.add_paragraph("Total")
add_tab_stop(p)
p.add_run("\\t1,500")
target = doc.add_paragraph("Bookmarked target")
add_bookmark(target, "target_here")
add_internal_link(doc.add_paragraph(), "jump", "target_here")
add_section(doc, "continuous", columns=2)
for i in range(4):
    doc.add_paragraph("双栏正文 two-column body text. " * 8)
add_column_break(doc.add_paragraph("右栏 Right column"))
add_section(doc, "continuous", columns=1)
add_section(doc, "new_page", orientation="landscape")
doc.add_paragraph("FY2025 landscape page")
assert replace_text(doc, "FY2025", "FY2026") == 1
doc.save("rich.docx")
"""


def test_helpers_build_rich_document(office_env, tmp_path):
    build(RICH_DOC, office_env, tmp_path)
    docx = tmp_path / "rich.docx"
    report = assert_checker_passes(
        docx, office_env, "年度报告", "数据来源：公司年报。", "FY2026 landscape page"
    )
    sections = report["summary"]["sections"]
    assert sections[0]["page_twips"] == {"w": "11906", "h": "16838"}
    assert (
        sections[-1]["page_twips"]["w"] == "16838"
        and sections[-1]["page_twips"].get("orient") == "landscape"
    )

    styles = part(docx, "word/styles.xml")
    default_fonts = styles.find(f"{W}docDefaults/{W}rPrDefault/{W}rPr/{W}rFonts")
    assert (
        default_fonts.get(f"{W}eastAsia") == "SimSun"
        and default_fonts.get(f"{W}ascii") == "Calibri"
    )
    assert not [n for n in styles.iter(f"{W}rFonts") if n.get(f"{W}asciiTheme")]
    heading = next(s for s in styles.iter(f"{W}style") if s.get(f"{W}styleId") == "Heading1")
    assert heading.find(f"{W}rPr/{W}rFonts").get(f"{W}eastAsia") == "SimHei"
    assert heading.find(f"{W}rPr/{W}sz").get(f"{W}val") == "32"

    body = part(docx, "word/document.xml").find(f"{W}body")
    num_ids = [n.get(f"{W}val") for n in body.iter(f"{W}numId")]
    assert len(set(num_ids)) == 3 and len(num_ids) == 5
    numbering = part(docx, "word/numbering.xml")
    kinds = [c.tag for c in numbering if c.tag in (f"{W}abstractNum", f"{W}num")]
    assert kinds.index(f"{W}num") > max(i for i, k in enumerate(kinds) if k == f"{W}abstractNum")
    assert not any("•" in text for text in paragraph_texts(docx))

    table = body.find(f"{W}tbl")
    assert [g.get(f"{W}w") for g in table.iter(f"{W}gridCol")] == ["4535", "2268"]
    assert table.find(f"{W}tr/{W}trPr/{W}tblHeader") is not None
    assert table.find(f".//{W}gridSpan").get(f"{W}val") == "2"
    assert any(s.get(f"{W}fill") == "FFF2CC" for s in table.iter(f"{W}shd"))
    last_row = table.findall(f"{W}tr")[-1]
    aligns = [c.find(f"{W}p/{W}pPr/{W}jc") for c in last_row.findall(f"{W}tc")]
    assert aligns[0].get(f"{W}val") == "left" and aligns[1].get(f"{W}val") == "right"
    assert any(b.get(f"{W}color") == "C00000" for b in table.iter(f"{W}bottom"))
    extent = part(docx, "word/document.xml").find(
        ".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"
    )
    assert int(extent.get("cx")) <= (11906 - 2 * 1134) * 635 + 635  # fits the text width

    rels = part(docx, "word/_rels/document.xml.rels")
    link = body.find(f".//{W}hyperlink[@{R_ID}]")
    target = next(r for r in rels if r.get("Id") == link.get(R_ID))
    assert (
        target.get("Target") == "https://example.com/filing"
        and target.get("TargetMode") == "External"
    )
    assert body.find(f".//{W}hyperlink[@{W}anchor='target_here']") is not None
    assert body.find(f".//{W}bookmarkStart[@{W}name='target_here']") is not None

    footnote_ref = body.find(f".//{W}footnoteReference")
    footnotes = part(docx, "word/footnotes.xml")
    note = next(n for n in footnotes if n.get(f"{W}id") == footnote_ref.get(f"{W}id"))
    assert "数据来源" in "".join(t.text or "" for t in note.iter(f"{W}t"))
    assert body.find(f".//{W}endnoteReference") is not None and "word/endnotes.xml" in members(docx)
    types = part(docx, "[Content_Types].xml")
    assert any(o.get("PartName") == "/word/footnotes.xml" for o in types)

    assert any(
        t.get(f"{W}leader") == "dot" and t.get(f"{W}val") == "right" for t in body.iter(f"{W}tab")
    )
    right = next(
        p for p in body.iter(f"{W}p") if "右栏" in "".join(t.text or "" for t in p.iter(f"{W}t"))
    )
    assert right.find(f"{W}r/{W}br").get(f"{W}type") == "column"  # break opens the paragraph
    cols = [c.get(f"{W}num") for c in body.iter(f"{W}cols")]
    assert "2" in cols and cols[-1] == "1"
    footer_names = [n for n in members(docx) if re.match(r"word/footer\d*\.xml", n)]
    footer = part(docx, footer_names[0])
    codes = " ".join(t.text or "" for t in footer.iter(f"{W}instrText"))
    assert "PAGE" in codes and "NUMPAGES" in codes


@pytest.mark.usefixtures("needs_dsoffice")
def test_rich_document_converts_and_renders(office_env, tmp_path):
    build(RICH_DOC, office_env, tmp_path)
    docx = tmp_path / "rich.docx"
    assert_converts_to_pdf(docx, office_env)
    result = run(
        ["dsoffice", "render", "--input", str(docx), "--output-dir", "pages", "--dpi", "60"],
        office_env,
        tmp_path,
    )
    manifest = json.loads(result.stdout)
    assert manifest["pageCount"] >= 2
    assert all(Path(image["path"]).is_file() for image in manifest["images"])


def test_replace_text_keeps_run_formatting(office_env, tmp_path):
    build(
        """
        from docx import Document
        from docx_helpers import replace_text
        doc = Document()
        p = doc.add_paragraph("Results for ")
        p.add_run("FY").bold = True
        p.add_run("2025").italic = True
        p.add_run(" and FY2025 again.")
        doc.sections[0].header.paragraphs[0].text = "Report FY2025"
        assert replace_text(doc, "FY2025", "FY2026") == 3
        doc.save("replaced.docx")
        """,
        office_env,
        tmp_path,
    )
    docx = tmp_path / "replaced.docx"
    assert paragraph_texts(docx)[0] == "Results for FY2026 and FY2026 again."
    runs = part(docx, "word/document.xml").find(f"{W}body/{W}p").findall(f"{W}r")
    assert runs[1].find(f"{W}t").text == "FY2026" and runs[1].find(f"{W}rPr/{W}b") is not None
    assert not runs[2].find(f"{W}t").text and runs[2].find(f"{W}rPr/{W}i") is not None
    header = next(n for n in members(docx) if n.startswith("word/header"))
    assert "Report FY2026" in "".join(t.text or "" for t in part(docx, header).iter(f"{W}t"))
    assert_checker_passes(docx, office_env, "FY2026")


# ---------------------------------------------------------------- table of contents

TOC_DOC = """
from docx import Document
from docx_helpers import *
doc = Document()
add_toc(doc, "1-2", title="Contents")
for chapter in range(1, 4):
    doc.add_page_break()
    doc.add_heading(f"Chapter {chapter}", 1)
    doc.add_paragraph("Body text. " * 40)
    doc.add_heading(f"Section {chapter}.1", 2)
    doc.add_heading(f"Too deep {chapter}", 3)
doc.save("draft.docx")
"""


@pytest.mark.usefixtures("needs_dsoffice")
def test_toc_is_filled_with_page_numbers(office_env, tmp_path):
    build(TOC_DOC, office_env, tmp_path)
    result = script("docx_toc.py", "draft.docx", "toc.docx", env=office_env, cwd=tmp_path)
    assert result.returncode == 0, result.stderr
    report = json.loads(result.stdout)
    entries = report["tocs"][0]["entries"]
    assert [e["text"] for e in entries] == [
        "Chapter 1",
        "Section 1.1",
        "Chapter 2",
        "Section 2.1",
        "Chapter 3",
        "Section 3.1",
    ]
    assert [e["page"] for e in entries] == ["2", "2", "3", "3", "4", "4"]
    assert report["page_numbers"].startswith("computed")

    docx = tmp_path / "toc.docx"
    assert_checker_passes(docx, office_env, "Chapter 3")
    body = part(docx, "word/document.xml").find(f"{W}body")
    codes = [t.text.strip() for t in body.iter(f"{W}instrText")]
    assert codes[0].startswith("TOC") and sum(c.startswith("PAGEREF _Toc") for c in codes) == 6
    bookmarks = {b.get(f"{W}name") for b in body.iter(f"{W}bookmarkStart")}
    assert {e["bookmark"] for e in entries} <= bookmarks
    assert_converts_to_pdf(docx, office_env)
    run(["dsoffice", "convert", "--input", "toc.docx", "--output", "toc.txt"], office_env, tmp_path)
    exported = (tmp_path / "toc.txt").read_text(encoding="utf-8-sig")
    toc_region = exported.split("\nChapter 1\n", 1)[0]  # LibreOffice shows the stored entries
    assert (
        "Chapter 2\t3" in toc_region
        and "Section 3.1\t4" in toc_region
        and "Too deep" not in toc_region
    )

    again = script("docx_toc.py", "toc.docx", "toc2.docx", env=office_env, cwd=tmp_path)
    assert again.returncode == 0, again.stderr
    assert [e["page"] for e in json.loads(again.stdout)["tocs"][0]["entries"]] == [
        "2",
        "2",
        "3",
        "3",
        "4",
        "4",
    ]
    body2 = part(tmp_path / "toc2.docx", "word/document.xml").find(f"{W}body")
    assert (
        len(
            [b for b in body2.iter(f"{W}bookmarkStart") if b.get(f"{W}name", "").startswith("_Toc")]
        )
        == 6
    )


# ---------------------------------------------------------------- review

CONTRACT = """
from docx import Document
from docx.shared import RGBColor
from docx_helpers import *
doc = Document()
doc.add_heading("Service Agreement 服务协议", 1)
p = doc.add_paragraph()
p.add_run("The Supplier shall deliver within ")
r = p.add_run("30 days"); r.bold = True; r.font.color.rgb = RGBColor(0xC0, 0, 0)
p.add_run(" of the order. Payment is due ")
p.add_run("in full").italic = True
p.add_run(" upon delivery.")
doc.add_paragraph("付款条件：买方应在收到发票后三十日内付款。")
doc.add_paragraph("This clause will be removed.")
p = doc.add_paragraph("See ")
add_hyperlink(p, "the annex", "https://example.com/annex")
p.add_run(" for prices.")
doc.add_paragraph("Governing law: Singapore. Governing law applies.")
doc.save("contract.docx")
"""


def _review_ops(tmp_path: Path) -> Path:
    ops = [
        {"op": "replace", "find": "30 days", "with": "45 days", "comment": "Extended per legal."},
        {"op": "replace", "find": "三十日", "with": "十五个工作日", "author": "韩梅梅"},
        {"op": "delete", "find": " in full"},
        {"op": "insert", "anchor": "the annex", "text": " and Schedule 2"},
        {"op": "delete-paragraph", "find": "This clause will be removed"},
        {
            "op": "insert-paragraph",
            "anchor": "Governing law",
            "text": "Disputes: arbitration 仲裁.",
        },
        {"op": "comment", "find": "Governing law", "text": "Second mention?", "occurrence": 2},
    ]
    path = tmp_path / "ops.json"
    path.write_text(json.dumps(ops, ensure_ascii=False), encoding="utf-8")
    return path


def test_review_writes_tracked_changes_and_comments(office_env, tmp_path):
    build(CONTRACT, office_env, tmp_path)
    ops = _review_ops(tmp_path)
    result = script(
        "docx_review.py",
        "apply",
        "contract.docx",
        "reviewed.docx",
        "--ops",
        ops,
        "--author",
        "Li Lei",
        "--date",
        DATE,
        env=office_env,
        cwd=tmp_path,
    )
    assert result.returncode == 0, result.stderr
    assert [r["changed"] for r in json.loads(result.stdout)["results"]] == [1] * 7
    docx = tmp_path / "reviewed.docx"
    assert_checker_passes(docx, office_env, "45 days")

    body = part(docx, "word/document.xml").find(f"{W}body")
    deletions = {"".join(t.text for t in d.iter(f"{W}delText")): d for d in body.iter(f"{W}del")}
    insertions = {"".join(t.text or "" for t in i.iter(f"{W}t")): i for i in body.iter(f"{W}ins")}
    assert (
        deletions["30 days"].get(f"{W}author") == "Li Lei"
        and deletions["30 days"].get(f"{W}date") == DATE
    )
    new_run = insertions["45 days"].find(f"{W}r/{W}rPr")
    assert (
        new_run.find(f"{W}b") is not None and new_run.find(f"{W}color").get(f"{W}val") == "C00000"
    )
    assert insertions["十五个工作日"].get(f"{W}author") == "韩梅梅"
    assert deletions[" in full"].find(f".//{W}i") is not None
    assert "This clause will be removed." in deletions
    assert " and Schedule 2" in insertions
    link = body.find(f".//{W}hyperlink")
    assert " and Schedule 2" not in "".join(t.text or "" for t in link.iter(f"{W}t"))
    ids = [n.get(f"{W}id") for n in body.iter() if n.tag in (f"{W}ins", f"{W}del")]
    assert len(ids) == len(set(ids))
    paragraph_marks = [p for p in body.iter(f"{W}p") if p.find(f"{W}pPr/{W}rPr/{W}del") is not None]
    assert len(paragraph_marks) == 1

    comments = part(docx, "word/comments.xml")
    texts = ["".join(t.text or "" for t in c.iter(f"{W}t")) for c in comments]
    assert texts == ["Extended per legal.", "Second mention?"]
    starts = [s.get(f"{W}id") for s in body.iter(f"{W}commentRangeStart")]
    ends = [s.get(f"{W}id") for s in body.iter(f"{W}commentRangeEnd")]
    refs = [s.get(f"{W}id") for s in body.iter(f"{W}commentReference")]
    assert sorted(starts) == sorted(ends) == sorted(refs) == ["0", "1"]
    types = part(docx, "[Content_Types].xml")
    assert any(o.get("PartName") == "/word/comments.xml" for o in types)
    rels = part(docx, "word/_rels/document.xml.rels")
    assert any(
        r.get("Type", "").endswith("/comments") and r.get("Target") == "comments.xml" for r in rels
    )

    listed = json.loads(
        script("docx_inspect.py", "comments", docx, "--json", env=office_env, cwd=tmp_path).stdout
    )
    assert [(c["anchor"], c["author"]) for c in listed] == [
        ("30 days45 days", "Li Lei"),
        ("Governing law", "Li Lei"),
    ]
    changes = json.loads(
        script("docx_inspect.py", "changes", docx, "--json", env=office_env, cwd=tmp_path).stdout
    )
    kinds = sorted(c["type"] for c in changes)
    assert kinds.count("delete") == 4 and kinds.count("insert") == 4
    assert "paragraph-mark-delete" in kinds and "paragraph-mark-insert" in kinds
    markup = script(
        "docx_inspect.py", "text", docx, "--view", "markup", env=office_env, cwd=tmp_path
    ).stdout
    assert "{-30 days-}{+45 days+}" in markup and "{-三十日-}{+十五个工作日+}" in markup


def test_accept_and_reject_review_edits(office_env, tmp_path):
    build(CONTRACT, office_env, tmp_path)
    ops = _review_ops(tmp_path)
    script(
        "docx_review.py",
        "apply",
        "contract.docx",
        "reviewed.docx",
        "--ops",
        ops,
        env=office_env,
        cwd=tmp_path,
    )
    for command in ("accept-all", "reject-all"):
        result = script(
            "docx_review.py",
            command,
            "reviewed.docx",
            f"{command}.docx",
            env=office_env,
            cwd=tmp_path,
        )
        assert result.returncode == 0, result.stderr
        out = tmp_path / f"{command}.docx"
        root = part(out, "word/document.xml")
        assert not [n for n in root.iter() if n.tag in (f"{W}ins", f"{W}del", f"{W}delText")]
        assert_checker_passes(out, office_env)
    accepted = paragraph_texts(tmp_path / "accept-all.docx")
    rejected = paragraph_texts(tmp_path / "reject-all.docx")
    assert (
        "The Supplier shall deliver within 45 days of the order. Payment is due upon delivery."
        in accepted
    )
    assert "付款条件：买方应在收到发票后十五个工作日内付款。" in accepted
    assert "See the annex and Schedule 2 for prices." in accepted
    assert "Disputes: arbitration 仲裁." in accepted
    assert not any("This clause" in t for t in accepted)
    assert (
        "The Supplier shall deliver within 30 days of the order."
        " Payment is due in full upon delivery." in rejected
    )
    assert "This clause will be removed." in rejected
    assert not any("Disputes" in t or "Schedule 2" in t for t in rejected)
    assert len(rejected) == len(paragraph_texts(tmp_path / "contract.docx"))
    # formatting of the replaced run survives acceptance
    body = part(tmp_path / "accept-all.docx", "word/document.xml").find(f"{W}body")
    run45 = next(
        r for r in body.iter(f"{W}r") if "".join(t.text or "" for t in r.iter(f"{W}t")) == "45 days"
    )
    assert run45.find(f"{W}rPr/{W}b") is not None


@pytest.mark.usefixtures("needs_dsoffice")
def test_review_outputs_convert_with_dsoffice(office_env, tmp_path):
    build(CONTRACT, office_env, tmp_path)
    ops = _review_ops(tmp_path)
    script(
        "docx_review.py",
        "apply",
        "contract.docx",
        "reviewed.docx",
        "--ops",
        ops,
        env=office_env,
        cwd=tmp_path,
    )
    script(
        "docx_review.py",
        "accept-all",
        "reviewed.docx",
        "accepted.docx",
        env=office_env,
        cwd=tmp_path,
    )
    script(
        "docx_review.py",
        "reject-all",
        "reviewed.docx",
        "rejected.docx",
        env=office_env,
        cwd=tmp_path,
    )
    for name in ("reviewed.docx", "accepted.docx", "rejected.docx"):
        assert_converts_to_pdf(tmp_path / name, office_env)
    run(
        ["dsoffice", "convert", "--input", "accepted.docx", "--output", "accepted.txt"],
        office_env,
        tmp_path,
    )
    exported = (tmp_path / "accepted.txt").read_text(encoding="utf-8-sig")
    assert "45 days" in exported and "30 days" not in exported and "This clause" not in exported


def test_review_refuses_bad_requests(office_env, tmp_path):
    build(CONTRACT, office_env, tmp_path)
    missing = script(
        "docx_review.py",
        "replace",
        "contract.docx",
        "out.docx",
        "--find",
        "nowhere",
        "--with",
        "x",
        env=office_env,
        cwd=tmp_path,
    )
    assert (
        missing.returncode == 1
        and "not found" in missing.stderr
        and not (tmp_path / "out.docx").exists()
    )
    script(
        "docx_review.py",
        "insert",
        "contract.docx",
        "one.docx",
        "--anchor",
        "Payment is due",
        "--text",
        " NEW",
        env=office_env,
        cwd=tmp_path,
    )
    nested = script(
        "docx_review.py",
        "delete",
        "one.docx",
        "two.docx",
        "--find",
        "NEW",
        env=office_env,
        cwd=tmp_path,
    )
    assert nested.returncode == 1 and "tracked insertion" in nested.stderr
    same = script(
        "docx_review.py",
        "accept-all",
        "contract.docx",
        "contract.docx",
        env=office_env,
        cwd=tmp_path,
    )
    assert same.returncode != 0


# ---------------------------------------------------------------- Word-style revisions

REV = f'w:author="Ann" w:date="{DATE}"'
WORD_REVISIONS = f"""
<w:p>
  <w:r><w:t xml:space="preserve">Keep </w:t></w:r>
  <w:ins w:id="101" {REV}><w:r><w:t>added</w:t></w:r></w:ins>
  <w:del w:id="102" {REV}><w:r><w:delText>removed</w:delText></w:r></w:del>
  <w:r><w:t xml:space="preserve"> end.</w:t></w:r>
</w:p>
<w:p>
  <w:r>
    <w:rPr><w:b/><w:rPrChange w:id="103" {REV}><w:rPr/></w:rPrChange></w:rPr>
    <w:t>Formatted</w:t>
  </w:r>
</w:p>
<w:p>
  <w:pPr>
    <w:jc w:val="center"/>
    <w:pPrChange w:id="104" {REV}><w:pPr><w:jc w:val="right"/></w:pPr></w:pPrChange>
  </w:pPr>
  <w:r><w:t>Aligned</w:t></w:r>
</w:p>
<w:p>
  <w:pPr><w:rPr><w:del w:id="105" {REV}/></w:rPr></w:pPr>
  <w:r><w:t>Joined</w:t></w:r>
</w:p>
<w:p><w:r><w:t>Tail</w:t></w:r></w:p>
<w:p>
  <w:pPr><w:rPr><w:ins w:id="106" {REV}/></w:rPr></w:pPr>
  <w:ins w:id="107" {REV}><w:r><w:t>New paragraph</w:t></w:r></w:ins>
</w:p>
<w:p>
  <w:moveFromRangeStart w:id="108" w:name="move1" {REV}/>
  <w:moveFrom w:id="109" {REV}><w:r><w:t>Moved</w:t></w:r></w:moveFrom>
  <w:moveFromRangeEnd w:id="108"/>
  <w:r><w:t xml:space="preserve"> source</w:t></w:r>
</w:p>
<w:p>
  <w:r><w:t xml:space="preserve">Target </w:t></w:r>
  <w:moveToRangeStart w:id="110" w:name="move1" {REV}/>
  <w:moveTo w:id="111" {REV}><w:r><w:t>Moved</w:t></w:r></w:moveTo>
  <w:moveToRangeEnd w:id="110"/>
</w:p>
<w:p>
  <w:ins w:id="112" {REV}>
    <w:r><w:t>Alpha</w:t></w:r>
    <w:del w:id="113" {REV}><w:r><w:delText>Beta</w:delText></w:r></w:del>
  </w:ins>
</w:p>
<w:tbl>
  <w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>
  <w:tblGrid><w:gridCol w:w="4000"/></w:tblGrid>
  <w:tr><w:tc><w:p><w:r><w:t>Row kept</w:t></w:r></w:p></w:tc></w:tr>
  <w:tr>
    <w:trPr><w:ins w:id="114" {REV}/></w:trPr>
    <w:tc><w:p><w:ins w:id="115" {REV}><w:r><w:t>Row added</w:t></w:r></w:ins></w:p></w:tc>
  </w:tr>
  <w:tr>
    <w:trPr><w:del w:id="116" {REV}/></w:trPr>
    <w:tc><w:p>
      <w:del w:id="117" {REV}><w:r><w:delText>Row dropped</w:delText></w:r></w:del>
    </w:p></w:tc>
  </w:tr>
</w:tbl>
<w:p/>
"""


def _word_revisions_doc(office_env, tmp_path) -> Path:
    build("from docx import Document\nDocument().save('blank.docx')\n", office_env, tmp_path)
    source = tmp_path / "blank.docx"
    target = tmp_path / "revisions.docx"
    with zipfile.ZipFile(source) as archive:
        items = [(info.filename, archive.read(info)) for info in archive.infolist()]
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in items:
            if name == "word/document.xml":
                text = data.decode("utf-8")
                start = text.index("<w:body>") + len("<w:body>")
                end = text.index("<w:sectPr")
                text = text[:start] + WORD_REVISIONS + text[end:]
                data = text.encode("utf-8")
            archive.writestr(name, data)
    return target


def test_accept_reject_resolve_word_revision_kinds(office_env, tmp_path):
    docx = _word_revisions_doc(office_env, tmp_path)
    assert_checker_passes(docx, office_env)
    changes = json.loads(
        script("docx_inspect.py", "changes", docx, "--json", env=office_env, cwd=tmp_path).stdout
    )
    kinds = {c["type"] for c in changes}
    assert kinds >= {
        "insert",
        "delete",
        "format",
        "paragraph-format",
        "paragraph-mark-delete",
        "paragraph-mark-insert",
        "move-from",
        "move-to",
        "row-insert",
        "row-delete",
    }
    assert {c["author"] for c in changes} == {"Ann"}

    script("docx_review.py", "accept-all", docx, "accepted.docx", env=office_env, cwd=tmp_path)
    script("docx_review.py", "reject-all", docx, "rejected.docx", env=office_env, cwd=tmp_path)
    accepted = [t for t in paragraph_texts(tmp_path / "accepted.docx") if t]
    rejected = [t for t in paragraph_texts(tmp_path / "rejected.docx") if t]
    assert accepted == [
        "Keep added end.",
        "Formatted",
        "Aligned",
        "JoinedTail",
        "New paragraph",
        " source",
        "Target Moved",
        "Alpha",
        "Row kept",
        "Row added",
    ]
    assert rejected == [
        "Keep removed end.",
        "Formatted",
        "Aligned",
        "Joined",
        "Tail",
        "Moved source",
        "Target ",
        "Row kept",
        "Row dropped",
    ]
    for name, bold, align in (("accepted.docx", True, "center"), ("rejected.docx", False, "right")):
        body = part(tmp_path / name, "word/document.xml").find(f"{W}body")
        formatted = next(
            p
            for p in body.iter(f"{W}p")
            if "Formatted" in "".join(t.text or "" for t in p.iter(f"{W}t"))
        )
        aligned = next(
            p
            for p in body.iter(f"{W}p")
            if "Aligned" in "".join(t.text or "" for t in p.iter(f"{W}t"))
        )
        assert (formatted.find(f".//{W}rPr/{W}b") is not None) is bold
        assert aligned.find(f"{W}pPr/{W}jc").get(f"{W}val") == align
        leftovers = [
            n.tag
            for n in body.iter()
            if n.tag.split("}")[1].endswith(("Change", "RangeStart", "RangeEnd"))
        ]
        assert leftovers == []
        assert_checker_passes(tmp_path / name, office_env)


@pytest.mark.usefixtures("needs_dsoffice")
def test_word_revision_results_convert(office_env, tmp_path):
    docx = _word_revisions_doc(office_env, tmp_path)
    script("docx_review.py", "accept-all", docx, "accepted.docx", env=office_env, cwd=tmp_path)
    script("docx_review.py", "reject-all", docx, "rejected.docx", env=office_env, cwd=tmp_path)
    for name in ("revisions.docx", "accepted.docx", "rejected.docx"):
        assert_converts_to_pdf(tmp_path / name, office_env)


# ---------------------------------------------------------------- reading


def test_inspect_text_outline(office_env, tmp_path):
    build(RICH_DOC, office_env, tmp_path)
    out = script("docx_inspect.py", "text", "rich.docx", env=office_env, cwd=tmp_path)
    assert out.returncode == 0, out.stderr
    lines = out.stdout.splitlines()
    assert "# 年度报告 Annual report" in lines
    assert "- 要点一 First" in lines and "  - Second level" in lines
    assert "1. Step one" in lines and "2. Step two" in lines and "1. Restarted" in lines
    assert "[table 1: 6 rows x 2 columns]" in lines and "| Region | Revenue |" in lines
    assert any(
        line.startswith("收入增长 12%[^1] see [the filing](https://example.com/filing)")
        for line in lines
    )
    assert "[^1] 数据来源：公司年报。" in lines
    assert any(
        line.startswith("[footer default, section 1]") and "{PAGE}" in line for line in lines
    )
    blocks = json.loads(
        script(
            "docx_inspect.py", "text", "rich.docx", "--json", env=office_env, cwd=tmp_path
        ).stdout
    )
    assert blocks[0]["heading_level"] == 1 and any(b["type"] == "table" for b in blocks)


# ---------------------------------------------------------------- unpack / pack


def test_unpack_edit_pack_roundtrip(office_env, tmp_path):
    build(CONTRACT, office_env, tmp_path)
    result = script("docx_pack.py", "unpack", "contract.docx", "x", env=office_env, cwd=tmp_path)
    assert result.returncode == 0, result.stderr
    document = tmp_path / "x/word/document.xml"
    xml = document.read_text(encoding="utf-8")
    assert "\n    <w:p>" in xml  # indented for editing
    document.write_text(xml.replace("for prices.", "for the price list."), encoding="utf-8")
    result = script("docx_pack.py", "pack", "x", "edited.docx", env=office_env, cwd=tmp_path)
    assert result.returncode == 0, result.stderr
    edited = tmp_path / "edited.docx"
    assert members(edited)[:3] == ["[Content_Types].xml", "_rels/.rels", "word/document.xml"]
    assert_checker_passes(edited, office_env, "for the price list.")
    with zipfile.ZipFile(edited) as archive:
        packed = archive.read("word/document.xml").decode("utf-8")
    assert (
        "\n    <w:p>" not in packed
        and '<w:t xml:space="preserve">The Supplier shall deliver within </w:t>' in packed
    )

    (tmp_path / "x/word/stray.bin").write_bytes(b"\0")
    stray = script("docx_pack.py", "pack", "x", "bad.docx", env=office_env, cwd=tmp_path)
    assert stray.returncode == 1 and "no content type" in stray.stderr
    (tmp_path / "x/word/stray.bin").unlink()
    document.write_text("<w:document", encoding="utf-8")
    broken = script("docx_pack.py", "pack", "x", "bad.docx", env=office_env, cwd=tmp_path)
    assert (
        broken.returncode == 1
        and "malformed XML" in broken.stderr
        and not (tmp_path / "bad.docx").exists()
    )
    again = script("docx_pack.py", "pack", "x", "edited.docx", env=office_env, cwd=tmp_path)
    assert again.returncode == 1 and "exists" in again.stderr


@pytest.mark.usefixtures("needs_dsoffice")
def test_packed_and_converted_documents(office_env, tmp_path):
    build(CONTRACT, office_env, tmp_path)
    script("docx_pack.py", "unpack", "contract.docx", "x", env=office_env, cwd=tmp_path)
    script("docx_pack.py", "pack", "x", "repacked.docx", env=office_env, cwd=tmp_path)
    assert_converts_to_pdf(tmp_path / "repacked.docx", office_env)
    # .odt stands in for legacy input: the same `dsoffice convert` path turns .doc into .docx
    run(
        ["dsoffice", "convert", "--input", "contract.docx", "--output", "contract.odt"],
        office_env,
        tmp_path,
    )
    run(
        ["dsoffice", "convert", "--input", "contract.odt", "--output", "back.docx"],
        office_env,
        tmp_path,
    )
    assert_checker_passes(tmp_path / "back.docx", office_env, "Governing law")
    text = script("docx_inspect.py", "text", "back.docx", env=office_env, cwd=tmp_path).stdout
    assert "# Service Agreement 服务协议" in text
