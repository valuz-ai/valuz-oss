#!/usr/bin/env python3
"""Audit a workbook against the financial-model conventions (reference/financial-models.md).

    valuz-python model_audit.py model.xlsx
    valuz-python model_audit.py model.xlsx --sheets Model DCF --axis both --strict

Read-only; the workbook is never modified. The JSON report on stdout covers:

  fonts                    font families used by non-empty cells (one family is expected)
  colour_coding            font colour per cell category: typed numbers (expected blue),
                           formulas (black), formulas reading other sheets (green), formulas
                           reading other workbooks (red). "observed" also reveals the scheme
                           an existing workbook already uses.
  hardcoded_numbers        numbers typed into formulas next to + - * / ^ (=B5*1.08), and
                           formulas built only from constants (=1250+380)
  inconsistent_formulas    a formula that breaks the pattern its neighbours share along a row
                           (the time axis); --axis column|both also checks down columns
  constants_in_formula_runs  a typed number sitting between formulas on that axis,
                           usually a pasted value that stopped the row from updating

Findings are prompts to review, not proof of a mistake. Exit 0, or 1 with --strict when
there are findings; 2 for invalid arguments; 3 when the workbook cannot be read.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import Token
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils.cell import column_index_from_string, get_column_letter
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

EXPECTED = {"input": "blue", "formula": "black", "cross_sheet": "green", "external": "red"}
# Excel theme colour index order: lt1, dk1, lt2, dk2, accent1-6, hlink, folHlink.
THEME_SLOTS = [
    "lt1",
    "dk1",
    "lt2",
    "dk2",
    "accent1",
    "accent2",
    "accent3",
    "accent4",
    "accent5",
    "accent6",
    "hlink",
    "folHlink",
]
OFFICE_THEME = [
    "FFFFFF",
    "000000",
    "E7E6E6",
    "44546A",
    "4472C4",
    "ED7D31",
    "A5A5A5",
    "FFC000",
    "5B9BD5",
    "70AD47",
    "0563C1",
    "954F72",
]
ARITHMETIC = {"+", "-", "*", "/", "^"}
CELLS = re.compile(r"\$?[A-Za-z]{1,3}\$?\d+(?::\$?[A-Za-z]{1,3}\$?\d+)?")
CELL = re.compile(r"(\$?)([A-Za-z]{1,3})(\$?)(\d+)")
COLUMNS = re.compile(r"(\$?)([A-Za-z]{1,3}):(\$?)([A-Za-z]{1,3})")
ROWS = re.compile(r"(\$?)(\d+):(\$?)(\d+)")


def local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def quote_sheet(name: str) -> str:
    if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_.]*", name):
        return name
    return "'" + name.replace("'", "''") + "'"


# ---------------------------------------------------------------------------
# Colours


def theme_colours(workbook) -> list[str]:
    """Theme colours by Excel theme index, read from the workbook's own theme when present."""
    data = getattr(workbook, "loaded_theme", None)
    if not data:
        return OFFICE_THEME
    try:
        root = ET.fromstring(data)
    except ET.ParseError:
        return OFFICE_THEME
    scheme = next((node for node in root.iter() if local(node.tag) == "clrScheme"), None)
    if scheme is None:
        return OFFICE_THEME
    found = {}
    for slot in scheme:
        for child in slot:
            value = child.get("val") if local(child.tag) == "srgbClr" else child.get("lastClr")
            if value:
                found[local(slot.tag)] = value.upper()
    return [found.get(name, OFFICE_THEME[index]) for index, name in enumerate(THEME_SLOTS)]


def colour_hex(color, theme: list[str]) -> str | None:
    """RRGGBB of a font colour; None when it cannot be resolved. Theme tints are ignored."""
    if color is None:
        return "000000"  # automatic
    kind = getattr(color, "type", None)
    if kind == "rgb":
        value = color.rgb
        return value[-6:].upper() if isinstance(value, str) and len(value) >= 6 else None
    if kind == "theme":
        index = color.theme
        return theme[index] if isinstance(index, int) and 0 <= index < len(theme) else None
    if kind == "indexed":
        index = color.indexed
        if index == 64:
            return "000000"  # system foreground
        if isinstance(index, int) and 0 <= index < len(COLOR_INDEX):
            return COLOR_INDEX[index][-6:].upper()
        return None
    if kind == "auto":
        return "000000"
    return None


def colour_name(value: str | None) -> str:
    if value is None:
        return "unknown"
    red, green, blue = (int(value[i : i + 2], 16) for i in (0, 2, 4))
    high, low = max(red, green, blue), min(red, green, blue)
    if high <= 0x40 or (high - low <= 0x20 and high <= 0x80):
        return "black"
    ranked = sorted((("red", red), ("green", green), ("blue", blue)), key=lambda pair: -pair[1])
    if ranked[0][1] - ranked[1][1] >= 0x30:
        return ranked[0][0]
    return "other"


# ---------------------------------------------------------------------------
# Formulas


def formula_of(cell) -> str | None:
    value = cell.value
    if isinstance(value, ArrayFormula):
        return value.text
    if isinstance(value, DataTableFormula):
        return None
    if cell.data_type == "f" and isinstance(value, str):
        return value
    return None


def tokens_of(formula: str) -> list[Token]:
    try:
        return [token for token in Tokenizer(formula).items if token.type != Token.WSPACE]
    except Exception:
        return []


def category(formula: str, host: str) -> str:
    """input/formula/cross_sheet/external for a formula cell on sheet ``host``."""
    other_sheet = False
    for token in tokens_of(formula):
        if token.type != Token.OPERAND or token.subtype != Token.RANGE or "!" not in token.value:
            continue
        prefix = token.value.rsplit("!", 1)[0].strip("'").replace("''", "'")
        if "[" in prefix:
            return "external"
        if prefix.lower() != host.lower():
            other_sheet = True
    return "cross_sheet" if other_sheet else "formula"


def typed_numbers(formula: str, allowed: set[float]) -> list[str]:
    """Numbers written into a formula where an assumption cell should be referenced."""
    tokens = tokens_of(formula)
    has_reference = any(t.type == Token.OPERAND and t.subtype == Token.RANGE for t in tokens)
    has_function = any(t.type == Token.FUNC for t in tokens)
    found = []
    for index, token in enumerate(tokens):
        if token.type != Token.OPERAND or token.subtype != Token.NUMBER:
            continue
        before = index - 1
        while before >= 0 and tokens[before].type == Token.OP_PRE:
            before -= 1
        after = index + 1
        percent = False
        while after < len(tokens) and tokens[after].type == Token.OP_POST:
            percent = percent or tokens[after].value == "%"
            after += 1
        beside_operator = any(
            0 <= position < len(tokens)
            and tokens[position].type == Token.OP_IN
            and tokens[position].value in ARITHMETIC
            for position in (before, after)
        )
        if not beside_operator and (has_reference or has_function):
            continue
        try:
            number = float(token.value) / (100 if percent else 1)
        except ValueError:
            continue
        if number not in allowed:
            found.append(token.value + ("%" if percent else ""))
    return found


def relative_reference(value: str, row: int, col: int) -> str:
    """Rewrite one range operand in R1C1 terms relative to (row, col); names pass through."""
    prefix = ""
    if "!" in value:
        prefix, value = value.rsplit("!", 1)
        prefix += "!"

    def row_part(dollar: str, number: str) -> str:
        return f"R{number}" if dollar else f"R[{int(number) - row}]"

    def col_part(dollar: str, letters: str) -> str:
        index = column_index_from_string(letters.upper())
        return f"C{index}" if dollar else f"C[{index - col}]"

    if CELLS.fullmatch(value):
        value = CELL.sub(lambda m: row_part(m[3], m[4]) + col_part(m[1], m[2]), value)
    elif match := COLUMNS.fullmatch(value):
        value = col_part(match[1], match[2]) + ":" + col_part(match[3], match[4])
    elif match := ROWS.fullmatch(value):
        value = row_part(match[1], match[2]) + ":" + row_part(match[3], match[4])
    return prefix + value


def signature(formula: str, row: int, col: int) -> str:
    parts = []
    for token in tokens_of(formula):
        if token.type == Token.OPERAND and token.subtype == Token.RANGE:
            parts.append(relative_reference(token.value, row, col))
        else:
            parts.append(token.value.upper() if token.type == Token.FUNC else token.value)
    return "".join(parts)


# ---------------------------------------------------------------------------
# Audit


def runs(entries: dict, axis: str):
    """Yield runs of adjacent formula/number cells along rows or columns."""
    lines: dict[int, list] = defaultdict(list)
    for (row, col), entry in entries.items():
        key, position = (row, col) if axis == "row" else (col, row)
        lines[key].append((position, row, col, entry))
    for items in lines.values():
        items.sort(key=lambda item: item[0])
        run: list = []
        for item in items:
            if run and item[0] != run[-1][0] + 1:
                yield run
                run = []
            run.append(item)
        if run:
            yield run


def audit_runs(entries: dict, axis: str, sheet: str, inconsistent: list, constants: list) -> None:
    for run in runs(entries, axis):
        formulas = [item for item in run if item[3][0] == "f"]
        if len(formulas) < 2:
            continue
        positions = [item[0] for item in formulas]
        for item in run:
            if item[3][0] == "n" and positions[0] < item[0] < positions[-1]:
                constants.append(
                    {
                        "cell": f"{quote_sheet(sheet)}!{get_column_letter(item[2])}{item[1]}",
                        "value": item[3][1],
                        "axis": axis,
                    }
                )
        if len(formulas) < 3:
            continue
        shapes = {item[0]: signature(item[3][1], item[1], item[2]) for item in formulas}
        pattern, count = Counter(shapes.values()).most_common(1)[0]
        if count < 2:
            continue
        # The first and last period often differ on purpose (opening link, terminal year),
        # so only a break with the pattern on both sides of it is reported.
        matching = [position for position, shape in shapes.items() if shape == pattern]
        inside = [item for item in formulas if matching[0] <= item[0] <= matching[-1]]
        outliers = [item for item in inside if shapes[item[0]] != pattern]
        if not outliers or len(outliers) > max(1, len(inside) // 4):
            continue
        example = next(item for item in formulas if shapes[item[0]] == pattern)
        for item in outliers:
            inconsistent.append(
                {
                    "cell": f"{quote_sheet(sheet)}!{get_column_letter(item[2])}{item[1]}",
                    "formula": item[3][1],
                    "axis": axis,
                    "pattern_cell": f"{get_column_letter(example[2])}{example[1]}",
                    "pattern_formula": example[3][1],
                }
            )


def audit(path: Path, sheets: list[str] | None, axes: list[str], allowed: set[float]) -> dict:
    workbook = load_workbook(path, data_only=False)
    names = sheets or workbook.sheetnames
    missing = [name for name in names if name not in workbook.sheetnames]
    if missing:
        raise KeyError(f"no such sheet: {', '.join(missing)}")
    theme = theme_colours(workbook)
    fonts: Counter = Counter()
    font_cells: dict[str, list[str]] = defaultdict(list)
    observed: dict[str, Counter] = {key: Counter() for key in EXPECTED}
    mismatches: list[dict] = []
    hardcoded: list[dict] = []
    inconsistent: list[dict] = []
    constants: list[dict] = []
    for name in names:
        sheet = workbook[name]
        entries: dict[tuple[int, int], tuple[str, object]] = {}
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                where = f"{quote_sheet(name)}!{cell.coordinate}"
                family = cell.font.name or "(default)"
                fonts[family] += 1
                font_cells[family].append(where)
                formula = formula_of(cell)
                if formula is not None:
                    kind = category(formula, name)
                    entries[(cell.row, cell.column)] = ("f", formula)
                    numbers = typed_numbers(formula, allowed)
                    if numbers:
                        hardcoded.append({"cell": where, "formula": formula, "numbers": numbers})
                elif isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    kind = "input"
                    entries[(cell.row, cell.column)] = ("n", cell.value)
                else:
                    continue
                colour = colour_name(colour_hex(cell.font.color, theme))
                observed[kind][colour] += 1
                if colour != EXPECTED[kind]:
                    mismatches.append(
                        {
                            "cell": where,
                            "category": kind,
                            "colour": colour,
                            "expected": EXPECTED[kind],
                        }
                    )
        for axis in axes:
            audit_runs(entries, axis, name, inconsistent, constants)
    main_font = fonts.most_common(1)[0][0] if fonts else None
    font_outliers = [
        cell for family, cells in font_cells.items() if family != main_font for cell in cells
    ]
    return {
        "sheets": names,
        "fonts": {"families": dict(fonts.most_common()), "consistent": len(fonts) <= 1},
        "font_outliers": font_outliers,
        "colour_coding": {
            "expected": EXPECTED,
            "observed": {key: dict(counter.most_common()) for key, counter in observed.items()},
        },
        "colour_mismatches": mismatches,
        "hardcoded_numbers": hardcoded,
        "inconsistent_formulas": inconsistent,
        "constants_in_formula_runs": constants,
    }


def summarise(result: dict, limit: int) -> dict:
    def capped(values: list) -> dict:
        return {"count": len(values), "cells": values[:limit], "truncated": len(values) > limit}

    notes = []
    for kind, counts in result["colour_coding"]["observed"].items():
        total = sum(counts.values())
        if total < 5:
            continue
        colour, count = max(counts.items(), key=lambda pair: pair[1])
        if colour != EXPECTED[kind] and count / total >= 0.8:
            notes.append(
                f"{kind} cells are mostly {colour} ({count}/{total}): the workbook has its own "
                "scheme; keep it when editing rather than recolouring."
            )
    report = {
        "sheets": result["sheets"],
        "fonts": {**result["fonts"], "outliers": capped(result["font_outliers"])},
        "colour_coding": {
            **result["colour_coding"],
            "mismatches": capped(result["colour_mismatches"]),
        },
        "hardcoded_numbers": capped(result["hardcoded_numbers"]),
        "inconsistent_formulas": capped(result["inconsistent_formulas"]),
        "constants_in_formula_runs": capped(result["constants_in_formula_runs"]),
    }
    report["finding_count"] = sum(
        len(result[key])
        for key in (
            "font_outliers",
            "colour_mismatches",
            "hardcoded_numbers",
            "inconsistent_formulas",
            "constants_in_formula_runs",
        )
    )
    report["notes"] = notes
    return report


def main() -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("input", type=Path, help="workbook to audit (.xlsx or .xlsm)")
    parser.add_argument("--sheets", nargs="+", help="audit only these sheets (exact names)")
    parser.add_argument(
        "--axis",
        choices=("row", "column", "both"),
        default="row",
        help="time axis for formula-pattern checks: row (periods across columns, default), "
        "column (periods down rows) or both",
    )
    parser.add_argument(
        "--allow-number",
        type=float,
        action="append",
        default=[],
        help="number that may appear inside formulas (0 and 1 are always allowed); repeatable",
    )
    parser.add_argument("--max-findings", type=int, default=50, help="cells listed per section")
    parser.add_argument("--strict", action="store_true", help="exit 1 when there are findings")
    parser.add_argument("--report", type=Path, help="also write the JSON report to this file")
    args = parser.parse_args()
    source = args.input.expanduser().resolve()
    if not source.is_file():
        parser.error(f"input not found: {source}")
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        parser.error("model_audit reads .xlsx or .xlsm files")
    axes = ["row", "column"] if args.axis == "both" else [args.axis]
    allowed = {0.0, 1.0, *args.allow_number}
    report: dict = {"file": str(source)}
    status = 0
    try:
        report.update(summarise(audit(source, args.sheets, axes, allowed), args.max_findings))
        if args.strict and report["finding_count"]:
            status = 1
    except KeyError as error:
        parser.error(str(error.args[0]))
    except (OSError, ValueError, zipfile.BadZipFile) as error:
        report["error"] = f"could not read workbook: {error}"
        status = 3
    text = json.dumps(report, ensure_ascii=False, indent=2) + "\n"
    if args.report is not None:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(text, encoding="utf-8")
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except (AttributeError, ValueError):
        pass
    sys.stdout.write(text)
    return status


if __name__ == "__main__":
    sys.exit(main())
