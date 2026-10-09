"""The bundled ``pptx`` skill: its scripts run through ``valuz-python`` and ``dsoffice``.

Fixture decks are built in ``tmp_path`` with python-pptx; every skill script is then
invoked exactly as a session would (``valuz-python <skill>/scripts/x.py``). Assertions
check real outcomes: slide order and package parts after duplicate/delete/move/arrange,
independent chart copies, outline content including notes, QA findings, thumbnail
grids, helper-built decks, the structural checker and a PDF conversion of the files the
scripts produced.
"""

from __future__ import annotations

import json
import posixpath
import re
import subprocess
import textwrap
import zipfile
from pathlib import Path

import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from tests.resources.office_skills.conftest import SKILLS_DIR, run

SKILL = SKILLS_DIR / "pptx"
SCRIPTS = SKILL / "scripts"
DEEPSEEK_CHECKER = (
    Path(__file__).resolve().parents[3]
    / "vendor/dsh-runtime/node_modules/@deepseek-ai/dsh-skill-office/assets/scripts/check_office.py"
)
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
P14_NS = "http://schemas.microsoft.com/office/powerpoint/2010/main"


# --------------------------------------------------------------------------- helpers


@pytest.fixture
def env(office_env: dict[str, str]) -> dict[str, str]:
    """office_env without bytecode writes: importing the helpers must not leave
    __pycache__ in the (normally read-only) skill directory."""
    return {**office_env, "PYTHONDONTWRITEBYTECODE": "1"}


def script(
    name: str, *args: object, env: dict[str, str], cwd: Path
) -> subprocess.CompletedProcess[str]:
    return run(["valuz-python", str(SCRIPTS / name), *map(str, args)], env, cwd)


def script_unchecked(
    name: str, *args: object, env: dict[str, str], cwd: Path
) -> subprocess.CompletedProcess[str]:
    argv = ["valuz-python", str(SCRIPTS / name), *map(str, args)]
    return subprocess.run(argv, env=env, cwd=cwd, capture_output=True, text=True)


def assert_checker_passes(deck: Path, env: dict[str, str], count: int | None = None) -> dict:
    args: list[object] = [deck]
    if count is not None:
        args += ["--count", count]
    report = json.loads(script("check_office.py", *args, env=env, cwd=deck.parent).stdout)
    assert report["verdict"] == "pass", report
    return report


def assert_converts_to_pdf(deck: Path, env: dict[str, str]) -> None:
    pdf = deck.with_suffix(".pdf")
    run(["dsoffice", "convert", "--input", str(deck), "--output", str(pdf)], env, deck.parent)
    assert pdf.read_bytes().startswith(b"%PDF")


def titles(deck: Path) -> list[str]:
    return [slide.shapes.title.text_frame.text for slide in Presentation(str(deck)).slides]


def widescreen() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


def titled_deck(path: Path, count: int) -> Path:
    prs = widescreen()
    for number in range(1, count + 1):
        prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = f"S{number}"
    prs.save(str(path))
    return path


def rich_deck(path: Path) -> Path:
    """Slide 2 has a picture, a chart and notes; slide 3 links to slide 2; slide 4 is hidden."""
    image = path.parent / "logo.png"
    Image.new("RGB", (240, 120), (30, 90, 200)).save(image)
    prs = widescreen()
    for number in range(1, 5):
        prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = f"S{number}"
    chart_slide = prs.slides[1]
    chart_slide.shapes.add_picture(str(image), Inches(0.6), Inches(1.8), Inches(3))
    data = CategoryChartData()
    data.categories = ["Q1", "Q2", "Q3"]
    data.add_series("Revenue", (10, 12, 15))
    chart_slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(4.5), Inches(1.8), Inches(7), Inches(4.5), data
    )
    chart_slide.notes_slide.notes_text_frame.text = "Talk about Q3"
    table = prs.slides[0].shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    for (r, c), text in {(0, 0): "Region", (0, 1): "Sales", (1, 0): "APAC", (1, 1): "42"}.items():
        table.cell(r, c).text = text
    link = prs.slides[2].shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    link.text_frame.text = "back to S2"
    link.click_action.target_slide = chart_slide
    prs.slides[3]._element.set("show", "0")
    ext_list = etree.SubElement(chart_slide._element, f"{{{P_NS}}}extLst")
    ext = etree.SubElement(ext_list, f"{{{P_NS}}}ext", uri="{BB962C8B-B14F-4D97-AF65-F5344CB8AC3E}")
    etree.SubElement(ext, f"{{{P14_NS}}}creationId", nsmap={"p14": P14_NS}, val="1234567")
    prs.save(str(path))
    return path


def package_xml(deck: Path) -> dict[str, etree._Element]:
    with zipfile.ZipFile(deck) as archive:
        return {
            name: etree.fromstring(archive.read(name))
            for name in archive.namelist()
            if name.endswith((".xml", ".rels"))
        }


def rel_targets(xml: dict[str, etree._Element], part: str) -> dict[str, tuple[str, str]]:
    rels = xml[posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")]
    out = {}
    for rel in rels:
        target = rel.get("Target")
        if rel.get("TargetMode") != "External":
            target = posixpath.normpath(posixpath.join(posixpath.dirname(part), target))
        out[rel.get("Id")] = (rel.get("Type").rsplit("/", 1)[-1], target)
    return out


def slide_parts(xml: dict[str, etree._Element]) -> list[str]:
    links = rel_targets(xml, "ppt/presentation.xml")
    ids = xml["ppt/presentation.xml"].find(f"{{{P_NS}}}sldIdLst")
    return [links[node.get(f"{{{R_NS}}}id")][1] for node in ids]


# --------------------------------------------------------------------------- skill package


def test_checker_is_the_unchanged_deepseek_copy() -> None:
    assert (SCRIPTS / "check_office.py").read_bytes() == DEEPSEEK_CHECKER.read_bytes()


def test_skill_md_frontmatter_links_and_environment() -> None:
    text = (SKILL / "SKILL.md").read_text(encoding="utf-8")
    front = re.match(r"---\nname: pptx\ndescription: (.+?)\n---\n", text, re.S)
    assert front and len(front.group(1)) > 80
    assert len(text.splitlines()) <= 300
    for link in re.findall(r"\]\((reference/[^)]+)\)", text):
        assert (SKILL / link).is_file(), link
    for name in re.findall(r"scripts/([a-z_]+\.py)", text):
        assert (SCRIPTS / name).is_file(), name
    every_doc = text + "".join(
        p.read_text(encoding="utf-8") for p in (SKILL / "reference").glob("*.md")
    )
    assert "../" not in every_doc
    for banned in (
        "load_workspace_dependencies",
        "present(",
        "read_image",
        "render_document",
        "<node>",
        "<cli>",
    ):
        assert banned not in every_doc, banned
    for required in ("valuz-python", "dsoffice", "deliver_artifacts"):
        assert required in text


# --------------------------------------------------------------------------- outline


def test_outline_reports_layout_titles_tables_charts_notes_and_hidden(tmp_path, env) -> None:
    deck = rich_deck(tmp_path / "deck.pptx")
    report = json.loads(
        script("outline.py", deck, "--json", "--layouts", env=env, cwd=tmp_path).stdout
    )
    assert report["slide_count"] == 4
    assert report["slide_size_in"] == [13.33, 7.5]
    first, second, _, fourth = report["slides"]
    assert first["layout"] == "Title Only" and first["title"] == "S1"
    table = next(s for s in first["shapes"] if s["kind"] == "table")
    assert table["rows"] == [["Region", "Sales"], ["APAC", "42"]]
    chart = next(s for s in second["shapes"] if s["kind"] == "chart")["chart"]
    assert chart["type"] == "COLUMN_CLUSTERED"
    assert chart["categories"] == ["Q1", "Q2", "Q3"]
    assert chart["series"] == [{"name": "Revenue", "values": [10.0, 12.0, 15.0]}]
    assert any(s["kind"] == "picture" for s in second["shapes"])
    assert second["notes"] == "Talk about Q3"
    assert fourth["hidden"] is True and first["hidden"] is False
    blank = next(layout for layout in report["layouts"] if layout["name"] == "Title Slide")
    assert [(ph["idx"], ph["type"]) for ph in blank["placeholders"]][:2] == [
        (0, "CENTER_TITLE"),
        (1, "SUBTITLE"),
    ]

    text = script("outline.py", deck, "--slides", "2", env=env, cwd=tmp_path).stdout
    assert "## Slide 2" in text and "## Slide 1 " not in text
    assert "series 'Revenue': 10, 12, 15" in text
    assert "Notes: Talk about Q3" in text


# --------------------------------------------------------------------------- slides.py


def test_duplicate_copies_charts_notes_and_shares_media(tmp_path, env, needs_dsoffice) -> None:
    deck = rich_deck(tmp_path / "deck.pptx")
    out = tmp_path / "dup.pptx"
    result = json.loads(
        script(
            "slides.py",
            "duplicate",
            deck,
            out,
            "--slide",
            2,
            "--copies",
            2,
            env=env,
            cwd=tmp_path,
        ).stdout
    )
    assert result["sources"] == [1, 2, None, None, 3, 4]
    assert titles(out) == ["S1", "S2", "S2", "S2", "S3", "S4"]

    xml = package_xml(out)
    names = set(xml)
    with zipfile.ZipFile(out) as archive:
        members = archive.namelist()
    assert len([m for m in members if m.startswith("ppt/charts/chart")]) == 3
    assert len([m for m in members if m.startswith("ppt/embeddings/")]) == 3
    assert len([m for m in members if m.startswith("ppt/media/")]) == 1
    order = slide_parts(xml)
    charts, images = set(), set()
    for slide in order[1:4]:
        links = rel_targets(xml, slide)
        kinds = {kind: target for kind, target in links.values()}
        charts.add(kinds["chart"])
        images.add(kinds["image"])
        notes = kinds["notesSlide"]
        back = {kind: target for kind, target in rel_targets(xml, notes).values()}
        assert back["slide"] == slide  # each notes slide points at its own slide
        assert (
            f"{posixpath.dirname(kinds['chart'])}/_rels/{posixpath.basename(kinds['chart'])}.rels"
            in names
        )
        # every relationship id used in the slide XML exists in its rels
        used = {
            v
            for node in xml[slide].iter()
            for k, v in node.attrib.items()
            if k.startswith(f"{{{R_NS}}}")
        }
        assert used <= set(links)
    assert len(charts) == 3 and len(images) == 1
    creation_ids = [
        node.get("val")
        for slide in order[1:4]
        for node in xml[slide].iter(f"{{{P14_NS}}}creationId")
    ]
    assert (
        len(creation_ids) == 3 and len(set(creation_ids)) == 3
    )  # copies get their own creation ids

    # the copy's chart is independent: changing it leaves the original untouched
    prs = Presentation(str(out))
    copy_chart = next(s for s in prs.slides[2].shapes if s.has_chart).chart
    new_data = CategoryChartData()
    new_data.categories = ["Q1", "Q2", "Q3"]
    new_data.add_series("Cost", (1, 2, 3))
    copy_chart.replace_data(new_data)
    edited = tmp_path / "edited.pptx"
    prs.save(str(edited))
    reread = Presentation(str(edited))
    original = next(s for s in reread.slides[1].shapes if s.has_chart).chart
    assert list(original.plots[0].series[0].values) == [10, 12, 15]
    assert reread.slides[2].notes_slide.notes_text_frame.text == "Talk about Q3"

    assert_checker_passes(out, env, count=6)
    assert_converts_to_pdf(out, env)

    to_end = tmp_path / "dup-end.pptx"
    script(
        "slides.py",
        "duplicate",
        deck,
        to_end,
        "--slide",
        1,
        "--to",
        5,
        env=env,
        cwd=tmp_path,
    )
    assert titles(to_end) == ["S1", "S2", "S3", "S4", "S1"]


def _add_sections_and_custom_show(path: Path) -> None:
    """Give the deck a section list and a custom show that both name every slide."""
    with zipfile.ZipFile(path) as archive:
        items = {name: archive.read(name) for name in archive.namelist()}
    root = etree.fromstring(items["ppt/presentation.xml"])
    ids = root.find(f"{{{P_NS}}}sldIdLst")
    show = etree.Element(f"{{{P_NS}}}custShowLst")
    custom = etree.SubElement(show, f"{{{P_NS}}}custShow", name="Short", id="0")
    sld_list = etree.SubElement(custom, f"{{{P_NS}}}sldLst")
    for node in ids:
        etree.SubElement(sld_list, f"{{{P_NS}}}sld").set(f"{{{R_NS}}}id", node.get(f"{{{R_NS}}}id"))
    root.find(f"{{{P_NS}}}defaultTextStyle").addprevious(show)
    ext_list = root.find(f"{{{P_NS}}}extLst")
    if ext_list is None:
        ext_list = etree.SubElement(root, f"{{{P_NS}}}extLst")
    ext = etree.SubElement(ext_list, f"{{{P_NS}}}ext", uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}")
    sections = etree.SubElement(ext, f"{{{P14_NS}}}sectionLst", nsmap={"p14": P14_NS})
    section = etree.SubElement(
        sections, f"{{{P14_NS}}}section", name="All", id="{6C3C2D50-2F7B-4E1B-9C35-0C9C1D1D2A01}"
    )
    section_ids = etree.SubElement(section, f"{{{P14_NS}}}sldIdLst")
    for node in ids:
        etree.SubElement(section_ids, f"{{{P14_NS}}}sldId", id=node.get("id"))
    items["ppt/presentation.xml"] = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in items.items():
            archive.writestr(name, data)


def test_delete_cleans_links_sections_and_custom_shows(tmp_path, env, needs_dsoffice) -> None:
    deck = rich_deck(tmp_path / "deck.pptx")
    _add_sections_and_custom_show(deck)
    assert_checker_passes(deck, env, count=4)
    out = tmp_path / "deleted.pptx"
    result = json.loads(
        script("slides.py", "delete", deck, out, "--slides", 2, env=env, cwd=tmp_path).stdout
    )
    assert result["sources"] == [1, 3, 4]
    assert titles(out) == ["S1", "S3", "S4"]

    xml = package_xml(out)
    with zipfile.ZipFile(out) as archive:
        members = archive.namelist()
    assert not [
        m
        for m in members
        if m.startswith(("ppt/charts/", "ppt/embeddings/", "ppt/notesSlides/notesSlide"))
    ]
    assert not [
        m for m in members if m.startswith("ppt/media/")
    ]  # the picture was only on the deleted slide
    root = xml["ppt/presentation.xml"]
    remaining = {node.get("id") for node in root.find(f"{{{P_NS}}}sldIdLst")}
    assert {node.get("id") for node in root.iter(f"{{{P14_NS}}}sldId")} == remaining
    pres_links = rel_targets(xml, "ppt/presentation.xml")
    shown = [node.get(f"{{{R_NS}}}id") for node in root.iter(f"{{{P_NS}}}sld")]
    assert len(shown) == 3 and all(rid in pres_links for rid in shown)
    linker = slide_parts(xml)[1]
    assert not list(xml[linker].iter(f"{{{A_NS}}}hlinkClick"))
    assert "slide" not in {kind for kind, _ in rel_targets(xml, linker).values()}

    assert_checker_passes(out, env, count=3)
    assert_converts_to_pdf(out, env)


def test_move_and_arrange_reorder_duplicate_and_drop(tmp_path, env) -> None:
    deck = titled_deck(tmp_path / "deck.pptx", 4)
    moved = tmp_path / "moved.pptx"
    result = json.loads(
        script(
            "slides.py", "move", deck, moved, "--slide", 4, "--to", 1, env=env, cwd=tmp_path
        ).stdout
    )
    assert result["sources"] == [4, 1, 2, 3]
    assert titles(moved) == ["S4", "S1", "S2", "S3"]

    arranged = tmp_path / "arranged.pptx"
    result = json.loads(
        script(
            "slides.py",
            "arrange",
            deck,
            arranged,
            "--order",
            "3,1,3,2-1",
            env=env,
            cwd=tmp_path,
        ).stdout
    )
    assert titles(arranged) == ["S3", "S1", "S3", "S2", "S1"]
    assert result["sources"] == [3, 1, None, 2, None]
    assert_checker_passes(arranged, env, count=5)
    assert len({p for p in package_xml(arranged) if p.startswith("ppt/slides/slide")}) == 5

    bad = script_unchecked(
        "slides.py",
        "arrange",
        deck,
        tmp_path / "bad.pptx",
        "--order",
        "1,9",
        env=env,
        cwd=tmp_path,
    )
    assert bad.returncode == 1 and json.loads(bad.stdout)["status"] == "error"
    assert not (tmp_path / "bad.pptx").exists()
    every = script_unchecked(
        "slides.py",
        "delete",
        deck,
        tmp_path / "none.pptx",
        "--slides",
        "1-4",
        env=env,
        cwd=tmp_path,
    )
    assert every.returncode == 2
    same = script_unchecked(
        "slides.py", "move", deck, deck, "--slide", 1, "--to", 2, env=env, cwd=tmp_path
    )
    assert same.returncode == 2


def as_potx(deck: Path) -> Path:
    """The same package saved as a PowerPoint template (.potx)."""
    template = deck.with_suffix(".potx")
    with (
        zipfile.ZipFile(deck) as source,
        zipfile.ZipFile(template, "w", zipfile.ZIP_DEFLATED) as out,
    ):
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"presentationml.presentation.main+xml", b"presentationml.template.main+xml"
                )
            out.writestr(item, data)
    return template


def test_potx_templates_are_accepted(tmp_path, env, needs_dsoffice) -> None:
    template = as_potx(titled_deck(tmp_path / "corporate.pptx", 3))
    outline = json.loads(script("outline.py", template, "--json", env=env, cwd=tmp_path).stdout)
    assert [slide["title"] for slide in outline["slides"]] == ["S1", "S2", "S3"]
    out = tmp_path / "from-template.pptx"
    script("slides.py", "arrange", template, out, "--order", "3,1", env=env, cwd=tmp_path)
    assert titles(out) == ["S3", "S1"]
    assert_checker_passes(out, env, count=2)
    qa = json.loads(script("content_qa.py", template, "--json", env=env, cwd=tmp_path).stdout)
    assert qa["slides"] == 3
    grid = json.loads(
        script(
            "thumbnails.py", template, "--output-dir", tmp_path / "t", env=env, cwd=tmp_path
        ).stdout
    )
    assert sorted(grid["slide_images"]) == ["1", "2", "3"]
    builder = tmp_path / "open.py"
    builder.write_text(
        textwrap.dedent(
            f"""
            import sys
            sys.path.insert(0, {str(SCRIPTS)!r})
            from pptx_helpers import open_presentation
            prs = open_presentation({str(template)!r})
            prs.slides[0].shapes.title.text_frame.text = "Filled"
            prs.save("filled-template.pptx")
            """
        ),
        encoding="utf-8",
    )
    run(["valuz-python", str(builder)], env, tmp_path)
    assert titles(tmp_path / "filled-template.pptx")[0] == "Filled"
    assert_checker_passes(tmp_path / "filled-template.pptx", env, count=3)
    assert_converts_to_pdf(tmp_path / "filled-template.pptx", env)


# --------------------------------------------------------------------------- content_qa.py


def qa_deck(path: Path) -> Path:
    prs = widescreen()
    clean = prs.slides.add_slide(prs.slide_layouts[5])
    clean.shapes.title.text = "Clean slide"
    box = clean.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    box.text_frame.word_wrap = True
    box.text_frame.text = "Short body text."

    leftover = prs.slides.add_slide(prs.slide_layouts[1])  # body placeholder left empty
    leftover.shapes.title.text = "Click to add title"
    note = leftover.shapes.add_textbox(Inches(1), Inches(6), Inches(6), Inches(0.5))
    note.text_frame.word_wrap = True
    note.text_frame.text = "Lorem ipsum dolor sit amet"

    crowded = prs.slides.add_slide(prs.slide_layouts[6])
    box = crowded.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.word_wrap = True
    box.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    box.text_frame.text = (
        "This long paragraph cannot fit into a one inch box at eighteen points " * 2
    )
    cjk = crowded.shapes.add_textbox(Inches(6), Inches(1), Inches(3), Inches(0.8))
    cjk.text_frame.word_wrap = True
    cjk.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    cjk.text_frame.text = (
        "这是一段很长的中文文本，用来测试中文字符在固定大小的文本框里是否会溢出，明显超过了高度。"
    )

    styled = prs.slides.add_slide(prs.slide_layouts[6])
    nowrap = styled.shapes.add_textbox(Inches(8), Inches(1), Inches(3), Inches(0.6))
    nowrap.text_frame.text = (
        "A single line of text that is far too long to stay on this slide when wrap is off"
    )
    tiny = styled.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(0.5))
    tiny.text_frame.word_wrap = True
    run_ = tiny.text_frame.paragraphs[0].add_run()
    run_.text = "tiny pale text"
    run_.font.size = Pt(7)
    run_.font.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
    band = styled.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(5), Inches(4), Inches(1.5)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    dark = styled.shapes.add_textbox(Inches(1.2), Inches(5.2), Inches(3.6), Inches(1))
    dark.text_frame.word_wrap = True
    dark_run = dark.text_frame.paragraphs[0].add_run()
    dark_run.text = "dark on dark TODO"
    dark_run.font.color.rgb = RGBColor(0x11, 0x22, 0x55)

    tabled = prs.slides.add_slide(prs.slide_layouts[6])
    table = tabled.shapes.add_table(12, 3, Inches(1), Inches(3), Inches(6), Inches(3)).table
    for r in range(12):
        for c in range(3):
            table.cell(r, c).text = f"row {r} col {c} with a little more text to wrap around"
    edge = tabled.shapes.add_textbox(Inches(12), Inches(6.8), Inches(2), Inches(1))
    edge.text_frame.text = "off"

    nested = prs.slides.add_slide(prs.slide_layouts[1])
    nested.shapes.title.text = "Many points"
    nested.placeholders[1].text_frame.text = "\n".join(
        f"Point number {n} with some words" for n in range(18)
    )
    group = nested.shapes.add_group_shape()
    inner = group.shapes.add_textbox(Inches(1), Inches(6.6), Inches(3), Inches(0.4))
    inner.name = "Grouped text"
    inner.text_frame.word_wrap = True
    inner.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    inner.text_frame.text = "Grouped text that wraps over several lines in a narrow box"
    prs.save(str(path))
    return path


def test_content_qa_flags_leftovers_empty_placeholders_and_overflow(tmp_path, env) -> None:
    deck = qa_deck(tmp_path / "qa.pptx")
    done = script_unchecked("content_qa.py", deck, "--json", env=env, cwd=tmp_path)
    assert done.returncode == 1
    report = json.loads(done.stdout)
    found = {(i["slide"], i["severity"], i["code"], i["shape"]) for i in report["issues"]}
    assert not [i for i in report["issues"] if i["slide"] == 1]
    assert (2, "error", "leftover-text", "Title 1") in found
    assert (2, "error", "leftover-text", "TextBox 3") in found
    assert (2, "error", "empty-placeholder", "Content Placeholder 2") in found
    assert (3, "error", "overflow", "TextBox 1") in found
    assert (3, "error", "overflow", "TextBox 2") in found  # CJK text counts as wide glyphs
    assert (4, "error", "overflow", "TextBox 1") in found  # wrap off, runs past the slide edge
    assert (4, "warning", "tiny-text", "TextBox 2") in found
    assert (4, "warning", "low-contrast", "TextBox 2") in found
    assert (4, "warning", "low-contrast", "TextBox 4") in found  # dark text on the dark band
    assert (4, "warning", "placeholder-marker", "TextBox 4") in found
    assert (5, "error", "overflow", "Table 1") in found
    assert (5, "error", "off-slide", "TextBox 2") in found
    assert (5, "warning", "dense", "") in found
    assert (6, "error", "overflow", "Grouped text") in found  # group members are checked too
    shrink = [
        i for i in report["issues"] if i["slide"] == 6 and i["shape"] == "Content Placeholder 2"
    ]
    assert [i["severity"] for i in shrink] == ["warning"] and "shrink-on-overflow" in shrink[0][
        "message"
    ]
    assert report["errors"] == sum(1 for i in report["issues"] if i["severity"] == "error")

    text = script_unchecked("content_qa.py", deck, env=env, cwd=tmp_path).stdout
    assert "slide 2 [error] leftover-text" in text


def test_content_qa_strict_fails_on_warnings_only(tmp_path, env) -> None:
    prs = widescreen()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Plan"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    box.text_frame.word_wrap = True
    box.text_frame.text = "Budget TBD"
    deck = tmp_path / "warn.pptx"
    prs.save(str(deck))
    relaxed = script_unchecked("content_qa.py", deck, "--json", env=env, cwd=tmp_path)
    assert relaxed.returncode == 0
    assert json.loads(relaxed.stdout)["warnings"] == 1
    strict = script_unchecked("content_qa.py", deck, "--strict", env=env, cwd=tmp_path)
    assert strict.returncode == 1


# --------------------------------------------------------------------------- thumbnails.py


def test_thumbnails_render_labelled_grids(tmp_path, env, needs_dsoffice) -> None:
    deck = rich_deck(tmp_path / "deck.pptx")
    out = tmp_path / "thumbs"
    summary = json.loads(
        script(
            "thumbnails.py",
            deck,
            "--output-dir",
            out,
            "--cols",
            2,
            "--rows",
            1,
            "--cell-width",
            320,
            env=env,
            cwd=tmp_path,
        ).stdout
    )
    assert summary["status"] == "ok"
    assert [grid["slides"] for grid in summary["grids"]] == [[1, 2], [3, 4]]
    assert sorted(summary["slide_images"]) == ["1", "2", "3", "4"]
    for grid in summary["grids"]:
        with Image.open(grid["path"]) as image:
            assert image.width == 14 + 2 * (320 + 14)
            assert image.height > 180
    assert (out / "slides" / "manifest.json").is_file()

    again = script_unchecked("thumbnails.py", deck, "--output-dir", out, env=env, cwd=tmp_path)
    assert again.returncode == 2  # output directory must be new

    reused = json.loads(
        script(
            "thumbnails.py",
            deck,
            "--output-dir",
            tmp_path / "subset",
            "--from-render",
            out / "slides",
            "--slides",
            "2,4",
            env=env,
            cwd=tmp_path,
        ).stdout
    )
    assert reused["grids"][0]["slides"] == [2, 4]


# --------------------------------------------------------------------------- pptx_helpers.py

BUILD_WITH_HELPERS = """
import sys
sys.path.insert(0, {scripts!r})
from PIL import Image
from pptx_helpers import *

Image.new("RGB", (800, 400), (200, 80, 40)).save("wide.png")
theme = PALETTES["navy"]
prs = new_presentation(theme=theme)
cover = blank_slide(prs, theme)
add_box(cover, 0, 0, 0.18, 7.5, fill=theme.primary)
add_text(cover, 0.9, 2.3, 8, 1.4, "2025 年度经营回顾", theme=theme, size=40, bold=True,
         anchor="bottom")
add_text(cover, 0.9, 3.8, 8, 0.5, "FY2025 Business Review", theme=theme, size=18, color=theme.muted)

s = content_slide(prs, "Revenue grew 18% on APAC demand", theme, subtitle="收入同比增长 18%")
left, right = columns(theme.content_box(subtitle=True), 2, weights=[2, 1])
add_chart(s, "column", ["2023", "2024", "2025"], {{"Revenue": [138, 163, 192]}}, *left, theme=theme,
          data_labels=True, number_format="#,##0", legend=None)
add_bullets(s, *right, ["APAC +32%", ("Vietnam, Thailand", 1), "毛利率 41.5%"], theme=theme)
set_notes(s, "Source: management accounts")

s = content_slide(prs, "Product lines", theme)
top, bottom = rows(theme.content_box(), 2, weights=[3, 1])
rows_ = [["Line", "Revenue", "YoY"], ["Cobots", "1.95", "+9%"], ["AMR", "1.36", "+41%"],
         ["Total", "3.31", "+21%"]]
add_table(s, rows_,
          top[0], top[1], top[2], theme=theme, bold_last_row=True)
for (x, y, w, h), label in zip(columns(bottom, 3), ["1", "2", "3"]):
    add_box(s, x, y, 0.6, 0.6, fill=theme.primary, kind="oval", text=label, theme=theme,
            size=16, color="FFFFFF")
    add_image(s, "wide.png", x + 0.8, y, w - 0.8, h, mode="cover")

s = content_slide(prs, "Mix and trend", theme)
a, b, c = columns(theme.content_box(), 3)
add_chart(s, "pie", ["A", "B", "C"], {{"Share": [50, 30, 20]}}, *a, theme=theme,
          data_labels=True, number_format="0")
add_chart(s, "line", ["Q1", "Q2", "Q3"], {{"x": [1, 3, 2], "y": [2, 2, 4]}}, *b, theme=theme)
add_chart(s, "scatter", None, {{"pts": [(1, 2), (2, 3), (3, 5)]}}, *c, theme=theme)
add_image(s, "wide.png", c[0], 6.0, c[2], 0.9, mode="contain")
prs.save("helpers.pptx")

dark = new_presentation(theme=PALETTES["midnight"])
slide = blank_slide(dark, PALETTES["midnight"])
add_text(slide, 1, 1, 8, 1, "Dark deck", theme=PALETTES["midnight"], size=40)
dark.save("dark.pptx")
"""


def test_helpers_build_a_clean_widescreen_deck(tmp_path, env, needs_dsoffice) -> None:
    builder = tmp_path / "build.py"
    builder.write_text(
        textwrap.dedent(BUILD_WITH_HELPERS.format(scripts=str(SCRIPTS))), encoding="utf-8"
    )
    run(["valuz-python", str(builder)], env, tmp_path)
    deck = tmp_path / "helpers.pptx"
    assert_checker_passes(deck, env, count=4)
    qa = json.loads(script_unchecked("content_qa.py", deck, "--json", env=env, cwd=tmp_path).stdout)
    assert qa["errors"] == 0, qa["issues"]

    prs = Presentation(str(deck))
    assert (prs.slide_width, prs.slide_height) == (Inches(13.333), Inches(7.5))
    title_layout = next(layout for layout in prs.slide_layouts if layout.name == "Title Only")
    assert title_layout.placeholders[0].width > Inches(11)  # 4:3 layouts were stretched to 16:9
    assert prs.slides[1].shapes.title.text_frame.text == "Revenue grew 18% on APAC demand"
    assert prs.slides[1].notes_slide.notes_text_frame.text == "Source: management accounts"
    assert not [
        ph
        for slide in prs.slides
        for ph in slide.placeholders
        if not ph.has_text_frame or not ph.text_frame.text.strip()
    ]

    xml = package_xml(deck)
    theme_xml = next(root for name, root in xml.items() if name.startswith("ppt/theme/"))
    assert theme_xml.find(f".//{{{A_NS}}}accent1/{{{A_NS}}}srgbClr").get("val") == "1F3A5F"
    assert (
        theme_xml.find(f".//{{{A_NS}}}minorFont/{{{A_NS}}}ea").get("typeface") == "Microsoft YaHei"
    )
    slides = slide_parts(xml)
    bullets = xml[slides[1]]
    assert len(list(bullets.iter(f"{{{A_NS}}}buChar"))) == 3
    assert not [
        t.text
        for t in bullets.iter(f"{{{A_NS}}}t")
        if t.text and t.text.lstrip().startswith(("•", "-"))
    ]
    assert all(r.find(f"{{{A_NS}}}ea") is not None for r in bullets.iter(f"{{{A_NS}}}rPr"))
    table = xml[slides[2]]
    assert table.find(f".//{{{A_NS}}}tableStyleId").text == "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    pictures = [p for p in prs.slides[2].shapes if p.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 3
    for picture in pictures:  # "cover": centred crop on one axis, aspect ratio kept
        sides = (
            (picture.crop_left, picture.crop_right)
            if picture.crop_left
            else (picture.crop_top, picture.crop_bottom)
        )
        assert sides[0] > 0 and abs(sides[0] - sides[1]) < 1e-6
    first_chart = next(s for s in prs.slides[1].shapes if s.has_chart).chart
    assert first_chart.value_axis.minimum_scale == 0
    assert first_chart.plots[0].has_data_labels

    dark = Presentation(str(tmp_path / "dark.pptx"))
    clr_map = dark.slide_master._element.find(f"{{{P_NS}}}clrMap")
    assert (clr_map.get("bg1"), clr_map.get("tx1")) == ("dk1", "lt1")

    assert_converts_to_pdf(deck, env)
    thumbs = json.loads(
        script("thumbnails.py", deck, "--output-dir", tmp_path / "qa", env=env, cwd=tmp_path).stdout
    )
    assert len(thumbs["slide_images"]) == 4
    assert not (SCRIPTS / "__pycache__").exists()  # nothing written into the skill directory


TEMPLATE_EDITS = """
import sys, json
sys.path.insert(0, {scripts!r})
from PIL import Image
from pptx import Presentation
from pptx_helpers import set_text, replace_text, replace_image

Image.new("RGB", (300, 300), (10, 160, 90)).save("square.png")
prs = Presentation("template.pptx")
slide = prs.slides[0]
set_text(slide.shapes.title, "季度回顾 Quarterly review")
body = slide.placeholders[1]
set_text(body, ["First point", ("Detail", 1), "Second point"])
count = replace_text(prs, "{{{{company}}}}", "Starlan")
picture = next(s for s in slide.shapes if s.shape_type == 13)
replace_image(picture, "square.png")
prs.save("filled.pptx")
print(json.dumps({{"replaced": count}}))
"""


def test_template_text_helpers_keep_formatting(tmp_path, env, needs_dsoffice) -> None:
    prs = widescreen()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title_run = slide.shapes.title.text_frame.paragraphs[0].add_run()
    title_run.text = "Template title"
    title_run.font.bold = True
    title_run.font.size = Pt(40)
    title_run.font.color.rgb = RGBColor(0xAA, 0x11, 0x22)
    body = slide.placeholders[1].text_frame
    body.text = "Sample body"
    body.paragraphs[0].runs[0].font.italic = True
    box = slide.shapes.add_textbox(Inches(1), Inches(6), Inches(6), Inches(0.6))
    paragraph = box.text_frame.paragraphs[0]
    for text, bold in (("Prepared for {{com", True), ("pany}} by us", False)):
        part = paragraph.add_run()
        part.text = text
        part.font.bold = bold
    image = tmp_path / "wide.png"
    Image.new("RGB", (600, 200), (200, 30, 30)).save(image)
    slide.shapes.add_picture(str(image), Inches(8), Inches(2), Inches(3), Inches(1))
    prs.save(str(tmp_path / "template.pptx"))

    builder = tmp_path / "fill.py"
    builder.write_text(
        textwrap.dedent(TEMPLATE_EDITS.format(scripts=str(SCRIPTS))), encoding="utf-8"
    )
    result = json.loads(run(["valuz-python", str(builder)], env, tmp_path).stdout)
    assert result["replaced"] == 1

    filled = tmp_path / "filled.pptx"
    slide = Presentation(str(filled)).slides[0]
    run0 = slide.shapes.title.text_frame.paragraphs[0].runs[0]
    assert run0.text == "季度回顾 Quarterly review"
    assert run0.font.bold and run0.font.size == Pt(40) and str(run0.font.color.rgb) == "AA1122"
    paragraphs = slide.placeholders[1].text_frame.paragraphs
    assert [(p.text, p.level) for p in paragraphs] == [
        ("First point", 0),
        ("Detail", 1),
        ("Second point", 0),
    ]
    assert all(p.runs[0].font.italic for p in paragraphs)
    runs = (
        next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX)
        .text_frame.paragraphs[0]
        .runs
    )
    assert "".join(r.text for r in runs) == "Prepared for Starlan by us"
    assert runs[0].text == "Prepared for Starlan" and runs[0].font.bold
    picture = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
    assert picture.width == Inches(3) and picture.height == Inches(1)
    assert picture.image.size == (300, 300)
    assert abs(picture.crop_top - 1 / 3) < 0.01 and abs(picture.crop_bottom - 1 / 3) < 0.01

    qa = json.loads(
        script_unchecked("content_qa.py", filled, "--json", env=env, cwd=tmp_path).stdout
    )
    assert not [i for i in qa["issues"] if i["code"] in ("leftover-text", "placeholder-marker")]
    assert_checker_passes(filled, env, count=1)
    assert_converts_to_pdf(filled, env)


def test_palettes_meet_contrast_targets(tmp_path, env) -> None:
    probe = tmp_path / "contrast.py"
    probe.write_text(
        textwrap.dedent(
            f"""
            import json, sys
            sys.path.insert(0, {str(SCRIPTS)!r})
            from pptx_helpers import PALETTES
            from content_qa import contrast
            rgb = lambda h: tuple(int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
            out = {{}}
            for name, t in PALETTES.items():
                out[name] = min(
                    contrast(rgb(t.text), rgb(t.bg)),
                    contrast(rgb(t.muted), rgb(t.bg)),
                    contrast(rgb(t.muted), rgb(t.surface)),
                    contrast(rgb(t.text), rgb(t.surface)),
                    contrast(rgb(t.on_primary), rgb(t.primary)),
                    contrast(rgb(t.highlight), rgb(t.bg)),
                )
            print(json.dumps(out))
            """
        ),
        encoding="utf-8",
    )
    worst = json.loads(run(["valuz-python", str(probe)], env, tmp_path).stdout)
    assert worst and all(value >= 4.5 for value in worst.values()), worst
